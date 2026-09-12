"""Create/refresh all staff logins directly from the curated SSDE manpower sheet.

    python seed_manpower.py "SSDE_manpower (7).xlsx"

The sheet is already app-shaped: it carries the exact Emp Code, Primary Role, and (for
dual-role people) a Team Lead ID. This script:
  * keeps every Emp Code EXACTLY as given (so the app's auto-generated next id continues
    from the next free number per prefix),
  * gives dual-role users BOTH hats — their primary role (FO../TC..) plus the team-lead
    hat via `also_team_lead=True` + `tl_emp_code=<TL..>` (one account, two views),
  * fills every HR field from the sheet,
  * is idempotent: matches existing accounts by email and updates them, never duplicates,
    and never overwrites a password the user has already changed.

Default first-login password for staff: Ssd@2026 (they're forced to change it).
The admin (role=admin) uses ADMIN_PASSWORD from .env.
"""
import os
import re
import sys
from datetime import datetime, date

import openpyxl
from sqlalchemy import String as _SAString

from app.database import Base, SessionLocal, engine
from app import models
from app.security import hash_password
from app.config import get_settings

DEFAULT_STAFF_PW = "Ssd@2026"

# Column length limits (PostgreSQL enforces these; SQLite ignores them). We clamp every
# string field to its limit so a stray long value (e.g. blood_group "Don't know") can never
# abort the whole seed.
_STR_LIMITS = {c.name: c.type.length for c in models.User.__table__.columns
               if isinstance(c.type, _SAString) and c.type.length}
_BLOOD_RE = re.compile(r"^(A|B|AB|O)[+-]$")


def _clamp(user):
    # values are kept as-is; we only guard against any string longer than its column
    # limit (Postgres enforces lengths; SQLite doesn't).
    for field, limit in _STR_LIMITS.items():
        v = getattr(user, field, None)
        if isinstance(v, str) and len(v) > limit:
            setattr(user, field, v[:limit].strip())


def _s(v):
    return str(v).strip() if v not in (None, "") else None


def _date(v):
    if not v:
        return None
    if isinstance(v, (datetime, date)):
        return v if isinstance(v, date) and not isinstance(v, datetime) else v.date()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(str(v)[:10], fmt).date()
        except ValueError:
            continue
    return None


def _sync_columns():
    """Bring the DB schema up to date (adds any missing columns) before seeding."""
    try:
        from app.seed import _sync_columns as s
        s()
    except Exception:
        pass


def main(path):
    if not os.path.exists(path):
        print(f"ERROR: manpower file not found: {path}")
        return 1
    Base.metadata.create_all(bind=engine)
    _sync_columns()

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Manpower"] if "Manpower" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c is not None else "" for c in rows[0]]
    H = {h: i for i, h in enumerate(hdr)}

    def g(r, col):
        return _s(r[H[col]]) if col in H and H[col] < len(r) else None

    settings = get_settings()
    db = SessionLocal()
    created = updated = dual = 0
    per_role = {}
    try:
        for r in rows[1:]:
            if not any(r):
                continue
            email = g(r, "Email")
            emp = g(r, "Emp Code")
            role = (g(r, "Primary Role") or "staff").strip()
            if not email or not emp:
                continue

            u = db.query(models.User).filter(models.User.email == email).first()
            is_new = u is None
            if is_new:
                u = models.User(email=email, role=role)
                db.add(u)

            u.name = g(r, "Name") or u.name
            u.role = role
            u.emp_code = emp
            u.hr_ref = g(r, "HR Ref")
            u.phone = g(r, "Phone")
            u.designation = g(r, "Designation")
            loc = g(r, "Location")
            u.location = loc
            u.branch = g(r, "Branch")           # kept exactly as in the sheet
            u.gender = g(r, "Gender")
            u.dob = _date(r[H["DOB"]]) if "DOB" in H else None
            u.joining_date = _date(r[H["DOJ"]]) if "DOJ" in H else None
            u.blood_group = g(r, "Blood")
            u.marital_status = g(r, "Marital")
            u.emergency_contact = g(r, "Emergency No")
            u.emergency_name = g(r, "Emergency Name")
            u.emergency_relation = g(r, "Relation")
            u.aadhar_number = g(r, "Aadhaar")
            u.pan_number = g(r, "PAN")
            u.bank_holder = g(r, "Bank Holder")
            u.bank_account = g(r, "Account")
            u.ifsc_code = g(r, "IFSC")
            u.bank_name = g(r, "Bank")
            u.current_address = g(r, "Current Address")
            u.rent_own = g(r, "Rent/Own")
            u.is_active = True

            # dual role: primary FO/TC + team-lead hat (keep their TL id as-is)
            tlid = g(r, "Team Lead ID")
            roles_label = (g(r, "Roles") or "")
            if tlid and role != "teamlead":
                u.also_team_lead = True
                u.tl_emp_code = tlid
                dual += 1
            else:
                u.also_team_lead = False
                u.tl_emp_code = None

            # head-office manager flag
            u.ho_manager = ("head office manager" in roles_label.lower())

            # clamp any over-length strings so Postgres never rejects the batch
            _clamp(u)

            if is_new:
                u.employment_type = "Full-time"
                u.profile_completed = True     # sheet already carries full HR detail

            # password: only set for new accounts (never clobber a changed one)
            if is_new or not u.hashed_password:
                if role == "admin":
                    u.hashed_password = hash_password(settings.admin_password)
                    u.must_change_password = False
                else:
                    u.hashed_password = hash_password(DEFAULT_STAFF_PW)
                    u.must_change_password = True

            created += int(is_new)
            updated += int(not is_new)
            per_role[role] = per_role.get(role, 0) + 1
        db.commit()
    finally:
        db.close()

    print("=== Manpower seed complete ===")
    print(f"  created {created}, updated {updated}, dual-role (also team lead) {dual}")
    print("  per role: " + ", ".join(f"{k}={v}" for k, v in sorted(per_role.items())))
    print(f"  staff first-login password: {DEFAULT_STAFF_PW}  (admin uses ADMIN_PASSWORD from .env)")
    return 0


if __name__ == "__main__":
    p = sys.argv[1] if len(sys.argv) > 1 else "SSDE_manpower.xlsx"
    raise SystemExit(main(p))
