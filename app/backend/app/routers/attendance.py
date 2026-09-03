"""Attendance + live presence.

- Check-in (with GPS) on the first login of the day → marked present (late if past the window).
- A ~30s heartbeat keeps presence (active/idle/offline) and accumulates worked / active / idle time.
- Shift end 19:00: the client shows a "keep working / check out" prompt; if ignored the server
  auto-checks-out at 19:00 (+5 min grace).
- Dashboards: today's board, a monthly sheet, a per-person day detail (with their audit log),
  and an Excel export. Scope: admin/HR/head-office see everyone; managers see their branch;
  team leads see their team; everyone else sees only themselves.
"""
import io
from datetime import datetime, date as date_cls, time as time_cls, timedelta

from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from .. import models, audit, presence
from ..database import get_db
from ..deps import get_current_user, require_roles
from ..storage import save_photo, resolve as resolve_photo

router = APIRouter(prefix="/api/attendance", tags=["attendance"])

IST = presence.IST
VIEW_ALL = ("admin", "hr", "headoffice")          # see everyone's attendance
DOWNLOAD_ROLES = ("admin", "hr", "headoffice")    # who can export
GRACE_MIN = 5                                      # minutes after shift end before auto-checkout
IDLE_LOGOUT_MIN = 30                               # continuous inactivity (post shift-end / overtime) → end session


def _tracked(u: models.User) -> bool:
    """Everyone is tracked for attendance except admins (viewers, not subjects)."""
    return (u.role or "") != "admin"


def _shift(role: str) -> tuple[str, str, str]:
    """(shift_start, shift_end, late_after) as HH:MM. Frontline (FOS + callers) must check in by
    10:00; all other roles get a looser 10:30 window, after which it's a late check-in (still
    present). FOS start their shift earlier at 08:00."""
    if role == "fos":
        return "08:00", "19:00", "10:00"
    if role == "telecaller":
        return "09:00", "19:00", "10:00"
    return "09:00", "19:00", "10:30"


def _hm(s: str) -> time_cls:
    h, m = s.split(":")
    return time_cls(int(h), int(m))


def _ist_now() -> datetime:
    return datetime.now(IST)


def _ist_today() -> date_cls:
    return _ist_now().date()


def _day_window_utc(d: date_cls):
    start = datetime.combine(d, time_cls(0, 0), tzinfo=IST).astimezone(presence.timezone.utc)
    return start, start + timedelta(days=1)


def _shift_dt(d: date_cls, hhmm: str) -> datetime:
    return datetime.combine(d, _hm(hhmm), tzinfo=IST)


# ---------------------------------------------------------------- activity rollup

def _activity_for(db: Session, uid: int, d: date_cls) -> dict:
    start, end = _day_window_utc(d)
    calls = db.query(func.count(models.CallLog.id)).filter(
        models.CallLog.caller_id == uid, models.CallLog.created_at >= start,
        models.CallLog.created_at < end).scalar() or 0
    visits = db.query(func.count(models.Visit.id)).filter(
        models.Visit.officer_id == uid, models.Visit.created_at >= start,
        models.Visit.created_at < end).scalar() or 0
    call_coll = db.query(func.coalesce(func.sum(models.CallLog.ptp_amount), 0)).filter(
        models.CallLog.caller_id == uid, models.CallLog.disposition.in_(("PAYMENT", "PAID")),
        models.CallLog.created_at >= start, models.CallLog.created_at < end).scalar() or 0
    visit_coll = db.query(func.coalesce(func.sum(models.Visit.amount_collected), 0)).filter(
        models.Visit.officer_id == uid, models.Visit.created_at >= start,
        models.Visit.created_at < end).scalar() or 0
    return {"calls": int(calls), "visits": int(visits),
            "collected": float(call_coll or 0) + float(visit_coll or 0)}


def _team_activity(db: Session, tl_id: int, d: date_cls) -> dict:
    """Combined calls / visits / collected for everyone reporting to this team lead (their whole
    team), for the day. Team leads themselves don't call/visit, so this rolls up the team."""
    ids = [uid for (uid,) in db.query(models.User.id).filter(models.User.team_lead_id == tl_id).all()]
    if not ids:
        return {"calls": 0, "visits": 0, "collected": 0.0}
    start, end = _day_window_utc(d)
    calls = db.query(func.count(models.CallLog.id)).filter(
        models.CallLog.caller_id.in_(ids), models.CallLog.created_at >= start,
        models.CallLog.created_at < end).scalar() or 0
    visits = db.query(func.count(models.Visit.id)).filter(
        models.Visit.officer_id.in_(ids), models.Visit.created_at >= start,
        models.Visit.created_at < end).scalar() or 0
    call_coll = db.query(func.coalesce(func.sum(models.CallLog.ptp_amount), 0)).filter(
        models.CallLog.caller_id.in_(ids), models.CallLog.disposition.in_(("PAYMENT", "PAID")),
        models.CallLog.created_at >= start, models.CallLog.created_at < end).scalar() or 0
    visit_coll = db.query(func.coalesce(func.sum(models.Visit.amount_collected), 0)).filter(
        models.Visit.officer_id.in_(ids), models.Visit.created_at >= start,
        models.Visit.created_at < end).scalar() or 0
    return {"calls": int(calls), "visits": int(visits),
            "collected": float(call_coll or 0) + float(visit_coll or 0)}


# ---------------------------------------------------------------- scope

def _scope_users(db: Session, viewer: models.User):
    q = db.query(models.User).filter(models.User.role != "admin", models.User.is_active.is_(True))
    if viewer.role in VIEW_ALL:
        return q
    if viewer.role == "manager":
        from sqlalchemy import select
        ids = select(models.User.id).where(models.User.branch == viewer.branch)
        return q.filter(or_(models.User.branch == viewer.branch, models.User.id.in_(ids)))
    if viewer.role == "teamlead" or getattr(viewer, "also_team_lead", False):
        return q.filter(or_(models.User.team_lead_id == viewer.id, models.User.id == viewer.id))
    # everyone else: only themselves
    return q.filter(models.User.id == viewer.id)


def _can_view(viewer: models.User, target: models.User, db: Session) -> bool:
    if viewer.id == target.id:
        return True
    if viewer.role in VIEW_ALL:
        return True
    if viewer.role == "manager":
        return target.branch == viewer.branch
    if viewer.role == "teamlead" or getattr(viewer, "also_team_lead", False):
        return target.team_lead_id == viewer.id
    return False


# ---------------------------------------------------------------- leave lookup

def _on_leave(db: Session, uid: int, d: date_cls) -> bool:
    try:
        rows = db.query(models.Leave).filter(models.Leave.user_id == uid,
                                             models.Leave.status == "approved").all()
        for lv in rows:
            s = getattr(lv, "start_date", None)
            e = getattr(lv, "end_date", None) or s
            if s and s <= d <= (e or s):
                return True
    except Exception:
        pass
    return False


# ---------------------------------------------------------------- helpers

def _today_row(db: Session, uid: int):
    return db.query(models.Attendance).filter(
        models.Attendance.user_id == uid, models.Attendance.date == _ist_today()).first()


def _row_out(db: Session, u: models.User, a: models.Attendance | None, d: date_cls,
             with_activity: bool = True) -> dict:
    ss, se, _la = _shift(u.role)
    # Calls / visits / collected only make sense for the frontline and, for a team lead, as their
    # team's combined total. Everyone else has no field/calling activity — omit it entirely.
    is_tl = u.role == "teamlead" or getattr(u, "also_team_lead", False)
    show_activity = u.role in ("fos", "telecaller") or is_tl
    if with_activity and show_activity:
        act = _team_activity(db, u.id, d) if is_tl else _activity_for(db, u.id, d)
    else:
        act = {}
    out = {
        "user_id": u.id, "name": u.name, "emp_code": u.emp_code, "role": u.role,
        "branch": u.branch, "shift_start": ss, "shift_end": se,
        "status": "absent", "late": False,
        "check_in_at": None, "check_out_at": None,
        "check_in_lat": None, "check_in_lng": None, "check_in_photo": None,
        "worked_seconds": 0, "active_seconds": 0, "idle_seconds": 0,
        "overtime": False, "overtime_seconds": 0, "auto_checkout": False, "platform": None,
        "presence": presence.presence_dict(u),
        "show_activity": bool(show_activity),
        "team_total": bool(is_tl),           # activity numbers are the team's combined total
        **act,
    }
    # Worked time stops at the person's LAST activity signal (see _live_worked), so anyone
    # who never checks out — field agent or caller — isn't credited up to 'now'. Compute it
    # once; for FOS also reuse the same GPS ping for the live-location line below.
    _fs, _fe = _day_window_utc(d)
    work_end, fos_lp = _work_end(db, u.id, u, _fs, _fe, a)
    if a:
        out.update({
            "status": a.status, "late": bool(a.late),
            "check_in_at": a.check_in_at.isoformat() if a.check_in_at else None,
            "check_out_at": a.check_out_at.isoformat() if a.check_out_at else None,
            "check_in_lat": a.check_in_lat, "check_in_lng": a.check_in_lng,
            "check_in_photo": resolve_photo(a.check_in_photo) if a.check_in_photo else None,
            "check_out_lat": a.check_out_lat, "check_out_lng": a.check_out_lng,
            "worked_seconds": _live_worked(a, work_end),
            "active_seconds": a.active_seconds or 0,
            "idle_seconds": a.idle_seconds or 0, "overtime": bool(a.overtime),
            "overtime_seconds": _overtime_secs(a, a.check_out_at or work_end or presence.now_utc()),
            "auto_checkout": bool(a.auto_checkout), "platform": a.last_platform,
        })
        # Checked out (manual or 7pm auto) ⇒ not online, regardless of a stray open tab.
        if a.check_out_at:
            out["presence"] = {**out["presence"], "state": "offline", "idle_seconds": 0}
    elif d.weekday() == 6:                       # Sunday = weekly off
        out["status"] = "weekoff"
    elif _on_leave(db, u.id, d):
        out["status"] = "leave"
    elif d > _ist_today():
        out["status"] = "—"
    # FOS carry a LIVE location (their latest GPS ping that day) so attendance can link to
    # live tracking; others just have their check-in point. (Reuses the ping fetched above.)
    if u.role == "fos" and fos_lp:
        out["live_lat"], out["live_lng"] = fos_lp.latitude, fos_lp.longitude
        out["live_at"] = fos_lp.created_at.isoformat() if fos_lp.created_at else None
    return out


def _shift_end_utc(a: models.Attendance) -> datetime:
    return _shift_dt(a.date, a.shift_end or "19:00").astimezone(presence.timezone.utc)


def _overtime_secs(a: models.Attendance, end) -> int:
    """Overtime = from when 'Continue working' was pressed → end. 0 if never continued."""
    if not a.overtime_start_at:
        return 0
    ots = presence._aware(a.overtime_start_at)
    end = presence._aware(end)
    if not ots or not end or end <= ots:
        return 0
    return int((end - ots).total_seconds())


def _live_worked(a: models.Attendance, end_override=None) -> int:
    """Worked seconds for the day.

    Normal day: check-in → end, where end is check-out, else end_override (the person's
    LAST activity signal — FOS GPS ping / anyone's last heartbeat — so an open tab isn't
    credited to 'now'), else now.

    Overtime day (they pressed 'Continue working' after shift end): the gap between shift
    end and the moment they continued does NOT count. Worked = [check-in → shift end]
    + [continue-press → end]. This holds whether they continued immediately at 7pm or
    came back after the 7pm auto-checkout and resumed.
    """
    if not a.check_in_at:
        return 0
    ci = presence._aware(a.check_in_at)
    if a.check_out_at:
        end = a.check_out_at
    elif end_override is not None:
        end = end_override
    else:
        end = presence.now_utc()
    end = presence._aware(end)
    if a.overtime_start_at:
        ots = presence._aware(a.overtime_start_at)
        se = _shift_end_utc(a)
        base_end = min(se, ots)                       # worked up to shift end, before overtime
        base = max(0, (min(base_end, end) - ci).total_seconds())
        ot = max(0, (end - ots).total_seconds()) if end > ots else 0
        return int(base + ot)
    return max(0, int((end - ci).total_seconds()))


def _work_end(db: Session, uid: int, u: models.User, start, end, a=None):
    """Last activity timestamp within the day, used to stop the worked-time clock for
    someone who never checked out (so we never credit up to 'now'):
      - FOS: their latest GPS ping that day (the 'live' signal), AND
      - anyone: their last heartbeat (last_seen) if it falls in the day, plus the
        attendance row's last heartbeat-touch (updated_at) as a per-day floor.
    Returns (timestamp_or_None, latest_fos_ping_or_None). None → caller falls back to now."""
    cands = []
    lp = None
    if u.role == "fos":
        lp = (db.query(models.LocationPing)
              .filter(models.LocationPing.officer_id == uid,
                      models.LocationPing.created_at >= start,
                      models.LocationPing.created_at < end)
              .order_by(models.LocationPing.created_at.desc()).first())
        if lp and lp.created_at:
            cands.append(presence._aware(lp.created_at))
    ls = presence._aware(getattr(u, "last_seen", None))
    if ls and start <= ls < end:
        cands.append(ls)
    ua = presence._aware(getattr(a, "updated_at", None)) if a is not None else None
    if ua and start <= ua < end:
        cands.append(ua)
    return (max(cands) if cands else None), lp


# ================================================================ self endpoints

@router.get("/me/today")
def my_today(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    a = _today_row(db, user.id)
    ss, se, la = _shift(user.role)
    return {
        "tracked": _tracked(user),
        "needs_checkin": _tracked(user) and (a is None or a.check_in_at is None),
        "shift_start": ss, "shift_end": se, "late_after": la,
        "attendance": _row_out(db, user, a, _ist_today()),
        "server_time": _ist_now().isoformat(),
    }


def _apply_checkin(db: Session, user: models.User, lat, lng, platform: str, photo_ref=None):
    """Shared check-in: mark present (late per shift window unless HO Manager), record GPS and,
    for FOS, the GPS-tagged selfie. Returns (attendance, already_checked_in)."""
    today = _ist_today()
    a = _today_row(db, user.id)
    if a and a.check_in_at:
        return a, True
    ss, se, la = _shift(user.role)
    now = presence.now_utc()
    late = (_ist_now().time() > _hm(la)) and not getattr(user, "ho_manager", False)
    if a is None:
        a = models.Attendance(user_id=user.id, date=today)
        db.add(a)
    a.check_in_at = now
    a.check_in_lat = lat
    a.check_in_lng = lng
    a.status = "present"
    a.late = bool(late)
    a.source = "checkin"
    a.shift_start, a.shift_end = ss, se
    a.last_platform = (platform or "web")[:10]
    if photo_ref:
        a.check_in_photo = photo_ref
    presence.touch(db, user.id, a.last_platform, active=True)
    audit.record(db, user, "checkin", None, entity_type="attendance",
                 detail=f"Checked in{' (late)' if late else ''} at {_ist_now():%H:%M}"
                        + (" · photo" if photo_ref else ""))
    db.commit()
    db.refresh(a)
    return a, False


@router.post("/checkin")
def checkin(body: dict = Body(default={}), db: Session = Depends(get_db),
            user: models.User = Depends(get_current_user)):
    if not _tracked(user):
        raise HTTPException(status_code=400, detail="Admins are not tracked for attendance.")
    a, already = _apply_checkin(db, user, body.get("lat"), body.get("lng"), body.get("platform") or "web")
    return {"ok": True, "already": already, "late": bool(a.late),
            "attendance": _row_out(db, user, a, _ist_today())}


@router.post("/checkin/photo")
async def checkin_with_photo(file: UploadFile = File(...), lat: str | None = Form(None),
                             lng: str | None = Form(None), platform: str = Form("android"),
                             db: Session = Depends(get_db),
                             user: models.User = Depends(get_current_user)):
    """FOS check-in with a GPS-tagged selfie from the phone (proof of presence)."""
    if not _tracked(user):
        raise HTTPException(status_code=400, detail="Admins are not tracked for attendance.")
    ct = (file.content_type or "").lower()
    if not ct.startswith("image/"):
        raise HTTPException(status_code=400, detail="Please attach a photo")
    content = await file.read()
    if len(content) > 8 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Photo too large (max 8 MB)")
    ref = save_photo(content, file.filename or "checkin.jpg", ct, folder="attendance")

    def _f(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    a, already = _apply_checkin(db, user, _f(lat), _f(lng), platform, photo_ref=ref)
    return {"ok": True, "already": already, "late": bool(a.late),
            "attendance": _row_out(db, user, a, _ist_today())}


@router.post("/checkout")
def checkout(body: dict = Body(default={}), db: Session = Depends(get_db),
             user: models.User = Depends(get_current_user)):
    a = _today_row(db, user.id)
    if not a or not a.check_in_at:
        raise HTTPException(status_code=400, detail="You haven't checked in today.")
    if a.check_out_at:
        return {"ok": True, "already": True, "attendance": _row_out(db, user, a, _ist_today())}
    now = presence.now_utc()
    a.check_out_at = now
    a.check_out_lat = body.get("lat")
    a.check_out_lng = body.get("lng")
    a.worked_seconds = _live_worked(a)
    a.overtime_seconds = _overtime_secs(a, now)   # only counts if they pressed "Continue working"
    audit.record(db, user, "checkout", None, entity_type="attendance",
                 detail=f"Checked out at {_ist_now():%H:%M}")
    db.commit()
    return {"ok": True, "attendance": _row_out(db, user, a, _ist_today())}


@router.post("/overtime")
def continue_overtime(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """User pressed 'Continue working'. Starts (or resumes) overtime from NOW:
      - if still checked in: just flags overtime and stamps the start.
      - if the 7pm auto-checkout already closed the day AND they've returned: REOPEN the day
        (clear the auto check-out) and start overtime from now — the away gap isn't counted.
    Overtime and active status begin from this moment."""
    a = _today_row(db, user.id)
    if not a or not a.check_in_at:
        return {"ok": False}
    now = presence.now_utc()
    if a.check_out_at:
        # only resume an AUTO checkout (manual "Check out" means they're done for the day)
        if not a.auto_checkout:
            return {"ok": False, "closed": True}
        a.check_out_at = None
        a.auto_checkout = False
        a.source = "checkin"
    a.overtime = True
    if not a.overtime_start_at:
        a.overtime_start_at = now                 # first continue-press of the day sets the origin
    a.worked_seconds = _live_worked(a)
    presence.touch(db, user.id, user.last_platform or "web", active=True)  # they clicked → active now
    audit.record(db, user, "overtime", None, entity_type="attendance",
                 detail=f"Continued working (overtime) from {_ist_now():%H:%M}")
    db.commit()
    return {"ok": True, "attendance": _row_out(db, user, a, _ist_today())}


@router.post("/heartbeat")
def heartbeat(body: dict = Body(default={}), db: Session = Depends(get_db),
              user: models.User = Depends(get_current_user)):
    """~30s presence ping. Carries platform + whether the user is currently active (visible &
    recent input, or on a call). Accumulates worked/active/idle time on today's row and tells the
    client about the shift-end prompt / auto-logout."""
    platform = (body.get("platform") or user.last_platform or "web")[:10]
    is_active = bool(body.get("active"))
    now = presence.now_utc()
    prev_seen = presence._aware(user.last_seen)

    a = _today_row(db, user.id)
    prompt = False            # show the "Continue working / Check out" modal
    resumed = False           # the modal is a post-auto-checkout resume offer (they came back)
    final_logout = False      # truly end the app session (long inactivity)
    minutes_to_end = None
    idle_min = 10 ** 9
    last_act = presence._aware(user.last_active_at)
    if last_act:
        idle_min = (now - last_act).total_seconds() / 60.0
    IDLE_LIMIT = IDLE_LOGOUT_MIN

    if a and a.check_in_at and not a.check_out_at:
        # attribute the gap since the last heartbeat to active or idle (ignore long gaps)
        if prev_seen:
            delta = int((now - prev_seen).total_seconds())
            if 0 < delta <= 2 * presence.ONLINE_TIMEOUT:
                if is_active:
                    a.active_seconds = (a.active_seconds or 0) + delta
                else:
                    a.idle_seconds = (a.idle_seconds or 0) + delta
        a.worked_seconds = _live_worked(a)
        a.last_platform = platform
        se_dt = _shift_dt(a.date, a.shift_end or "19:00")
        minutes_to_end = int((se_dt - _ist_now()).total_seconds() // 60)
        if a.overtime:
            a.overtime_seconds = _overtime_secs(a, now)
            # continuous inactivity while in overtime → end overtime at last activity + log out
            if idle_min > IDLE_LIMIT and not is_active:
                end_ts = last_act or now
                a.check_out_at = end_ts
                a.auto_checkout = True
                a.source = "auto"
                a.worked_seconds = _live_worked(a)
                a.overtime_seconds = _overtime_secs(a, end_ts)
                final_logout = True
                audit.record(db, user, "checkout", None, entity_type="attendance",
                             detail="Auto checked-out (inactive during overtime)")
        else:
            if _ist_now() >= se_dt + timedelta(minutes=GRACE_MIN):
                # no response for 5 min after shift end → auto check-out AT shift end.
                # This does NOT log them out: they stay in the app (now offline) and can press
                # "Continue working" to resume overtime if they return.
                a.check_out_at = se_dt.astimezone(presence.timezone.utc)
                a.worked_seconds = _live_worked(a)
                a.auto_checkout = True
                a.source = "auto"
                audit.record(db, user, "checkout", None, entity_type="attendance",
                             detail="Auto checked-out at shift end (no response)")
            elif _ist_now() >= se_dt:
                prompt = True                    # first shift-end prompt

    elif a and a.check_in_at and a.check_out_at and a.auto_checkout:
        # day was auto-closed at shift end but the tab is still open
        if is_active:
            prompt = True                        # they came back → offer to resume overtime
            resumed = True
        elif idle_min > IDLE_LIMIT:
            final_logout = True                  # away long enough → actually end the session

    # A checked-out person is offline regardless of a still-open tab: skip the presence refresh
    # (so last_seen goes stale) AND force offline in the response.
    checked_out = bool(a and a.check_out_at)
    if not checked_out:
        presence.touch(db, user.id, platform, active=is_active)
    db.commit()
    pres = presence.presence_dict(db.query(models.User).get(user.id))
    if checked_out:
        pres = {**pres, "state": "offline", "idle_seconds": 0}
    return {
        "presence": pres,
        "prompt_overtime": prompt, "resumed": resumed,
        "auto_logout": final_logout,
        "minutes_to_shift_end": minutes_to_end,
        "checked_out": checked_out,
    }


# ================================================================ sweep (auto-checkout)

def sweep(db: Session) -> None:
    """Close out today's still-open rows:
      - not in overtime: auto check-out at shift end + grace (handles closed tabs).
      - in overtime: auto check-out at their last activity once they've been unseen /
        inactive beyond the idle limit (handles someone who left mid-overtime).
    Safe to call often; only touches today's un-checked-out rows."""
    today = _ist_today()
    now = presence.now_utc()
    rows = db.query(models.Attendance).filter(
        models.Attendance.date == today, models.Attendance.check_in_at.isnot(None),
        models.Attendance.check_out_at.is_(None)).all()
    changed = False
    for a in rows:
        se_dt = _shift_dt(a.date, a.shift_end or "19:00")
        if a.overtime:
            u = a.user
            last = presence._aware(getattr(u, "last_active_at", None)) or presence._aware(getattr(u, "last_seen", None))
            if last and (now - last).total_seconds() > IDLE_LOGOUT_MIN * 60:
                a.check_out_at = last
                a.auto_checkout = True
                a.source = "auto"
                a.worked_seconds = _live_worked(a)
                a.overtime_seconds = _overtime_secs(a, last)
                changed = True
        elif _ist_now() >= se_dt + timedelta(minutes=GRACE_MIN):
            a.check_out_at = se_dt.astimezone(presence.timezone.utc)
            a.worked_seconds = _live_worked(a)
            a.auto_checkout = True
            a.source = "auto"
            changed = True
    if changed:
        db.commit()


# ================================================================ dashboards

@router.get("/day")
def day_board(date: str | None = None, role: str | None = None, user_id: int | None = None,
              db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    d = date_cls.fromisoformat(date) if date else _ist_today()
    if d == _ist_today():
        sweep(db)
    users = _scope_users(db, user)
    if role:
        users = users.filter(models.User.role == role)
    if user_id:
        users = users.filter(models.User.id == user_id)
    users = users.order_by(models.User.name).all()
    ids = [u.id for u in users]
    amap = {}
    if ids:
        for a in db.query(models.Attendance).filter(
                models.Attendance.date == d, models.Attendance.user_id.in_(ids)).all():
            amap[a.user_id] = a
    rows = [_row_out(db, u, amap.get(u.id), d) for u in users]
    summary = {
        "present": sum(1 for r in rows if r["status"] in ("present",)),
        "late": sum(1 for r in rows if r["late"]),
        "absent": sum(1 for r in rows if r["status"] == "absent"),
        "leave": sum(1 for r in rows if r["status"] == "leave"),
        "online": sum(1 for r in rows if r["presence"]["state"] in ("active", "idle")),
        "total": len(rows),
    }
    return {"date": d.isoformat(), "rows": rows, "summary": summary,
            "can_download": user.role in DOWNLOAD_ROLES}


@router.get("/user/{uid}")
def user_detail(uid: int, date: str | None = None, db: Session = Depends(get_db),
                user: models.User = Depends(get_current_user)):
    target = db.query(models.User).get(uid)
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    if not _can_view(user, target, db):
        raise HTTPException(status_code=403, detail="Not allowed to view this employee's attendance.")
    d = date_cls.fromisoformat(date) if date else _ist_today()
    a = db.query(models.Attendance).filter(
        models.Attendance.user_id == uid, models.Attendance.date == d).first()
    start, end = _day_window_utc(d)
    logs = []
    try:
        alogs = db.query(models.AuditLog).filter(
            models.AuditLog.created_at >= start, models.AuditLog.created_at < end,
            or_(models.AuditLog.actor_id == uid, models.AuditLog.target_user_id == uid)
        ).order_by(models.AuditLog.created_at.desc()).limit(200).all()
        for lg in alogs:
            logs.append({"at": lg.created_at.isoformat() if lg.created_at else None,
                         "action": lg.action, "detail": lg.detail, "entity": lg.entity_type})
    except Exception:
        pass
    return {"date": d.isoformat(), "row": _row_out(db, target, a, d), "audit": logs}


@router.get("/user/{uid}/visits")
def user_visits(uid: int, date: str | None = None, db: Session = Depends(get_db),
                user: models.User = Depends(get_current_user)):
    """Field visits logged by this employee on the given day, each with its GPS-camera
    photo, geo-fence distance, disposition, note and amount — for the attendance drawer."""
    target = db.query(models.User).get(uid)
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    if not _can_view(user, target, db):
        raise HTTPException(status_code=403, detail="Not allowed to view this employee's visits.")
    d = date_cls.fromisoformat(date) if date else _ist_today()
    start, end = _day_window_utc(d)
    rows = (db.query(models.Visit)
            .filter(models.Visit.officer_id == uid,
                    models.Visit.created_at >= start, models.Visit.created_at < end)
            .order_by(models.Visit.created_at.desc()).all())
    out = []
    for v in rows:
        c = v.case
        label = None
        if c:
            ref = c.card_no or c.account_no or ""
            label = " · ".join(x for x in [c.customer_name, c.bank, ref] if x)
        out.append({
            "id": v.id, "at": v.created_at.isoformat() if v.created_at else None,
            "case_id": v.case_id, "case_label": label,
            "photo": resolve_photo(v.photo_path) if v.photo_path else None,
            "lat": v.latitude, "lng": v.longitude,
            "distance_m": round(v.distance_from_case_m, 0) if v.distance_from_case_m is not None else None,
            "location_correct": v.location_correct, "person_moved": bool(v.person_moved),
            "disposition": v.disposition, "note": v.note,
            "paid": bool(v.paid), "amount": float(v.amount_collected or 0),
        })
    return {"date": d.isoformat(), "user_id": uid, "name": target.name,
            "count": len(out), "visits": out}


@router.get("/user/{uid}/calls")
def user_calls(uid: int, date: str | None = None, db: Session = Depends(get_db),
               user: models.User = Depends(get_current_user)):
    """Call logs made by this employee on the given day — disposition, PTP, note and the
    case each call was on — for the attendance drawer."""
    target = db.query(models.User).get(uid)
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    if not _can_view(user, target, db):
        raise HTTPException(status_code=403, detail="Not allowed to view this employee's call logs.")
    d = date_cls.fromisoformat(date) if date else _ist_today()
    start, end = _day_window_utc(d)
    rows = (db.query(models.CallLog)
            .filter(models.CallLog.caller_id == uid,
                    models.CallLog.created_at >= start, models.CallLog.created_at < end)
            .order_by(models.CallLog.created_at.desc()).all())
    out = []
    for cl in rows:
        c = cl.case
        label = None
        if c:
            ref = c.card_no or c.account_no or ""
            label = " · ".join(x for x in [c.customer_name, c.bank, ref] if x)
        out.append({
            "id": cl.id, "at": cl.created_at.isoformat() if cl.created_at else None,
            "case_id": cl.case_id, "case_label": label,
            "disposition": cl.disposition, "note": cl.note,
            "ptp_amount": float(cl.ptp_amount or 0),
            "ptp_date": cl.ptp_date.isoformat() if cl.ptp_date else None,
        })
    return {"date": d.isoformat(), "user_id": uid, "name": target.name,
            "count": len(out), "calls": out}


@router.get("/month")
def month_sheet(month: str | None = None, role: str | None = None, user_id: int | None = None,
                db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Per-user month matrix + totals (days present, late, leave, absent, total hours)."""
    m = month or _ist_now().strftime("%Y-%m")
    yy, mm = int(m[:4]), int(m[5:7])
    first = date_cls(yy, mm, 1)
    ndays = (date_cls(yy + (mm == 12), (mm % 12) + 1, 1) - first).days
    days = [first + timedelta(days=i) for i in range(ndays)]
    today = _ist_today()

    users = _scope_users(db, user)
    if role:
        users = users.filter(models.User.role == role)
    if user_id:
        users = users.filter(models.User.id == user_id)
    users = users.order_by(models.User.name).all()
    ids = [u.id for u in users]

    arows: dict[int, dict[str, models.Attendance]] = {}
    if ids:
        q = db.query(models.Attendance).filter(
            models.Attendance.user_id.in_(ids), models.Attendance.date >= first,
            models.Attendance.date <= days[-1])
        for a in q.all():
            arows.setdefault(a.user_id, {})[a.date.isoformat()] = a

    people = []
    for u in users:
        per_day = {}
        present = late = leave = absent = weekoff = 0
        worked = 0
        for d in days:
            a = arows.get(u.id, {}).get(d.isoformat())
            if a and a.check_in_at:
                st = "L" if a.late else "P"       # Late still counts present
                present += 1
                late += 1 if a.late else 0
                # Anyone who never checked out that day: stop the clock at their last
                # activity (FOS last GPS ping / anyone's last heartbeat) instead of 'now'.
                eo = None
                if not a.check_out_at:
                    _fs, _fe = _day_window_utc(d)
                    eo, _ = _work_end(db, u.id, u, _fs, _fe, a)
                worked += _live_worked(a, eo)
            elif d.weekday() == 6:
                st = "W"; weekoff += 1
            elif _on_leave(db, u.id, d):
                st = "LV"; leave += 1
            elif d > today:
                st = ""
            else:
                st = "A"; absent += 1
            per_day[d.isoformat()] = st
        people.append({
            "user_id": u.id, "name": u.name, "emp_code": u.emp_code, "role": u.role,
            "branch": u.branch, "days": per_day,
            "present": present, "late": late, "leave": leave, "absent": absent,
            "weekoff": weekoff, "worked_hours": round(worked / 3600.0, 1),
        })
    return {"month": m, "days": [d.isoformat() for d in days], "people": people,
            "can_download": user.role in DOWNLOAD_ROLES}


@router.get("/presence")
def presence_snapshot(ids: str | None = None, db: Session = Depends(get_db),
                      user: models.User = Depends(get_current_user)):
    """Live presence for a set of users (for profile cards / team lists). `ids` = CSV of user ids;
    omitted = everyone in the viewer's scope."""
    if ids:
        idlist = [int(x) for x in ids.split(",") if x.strip().isdigit()]
        us = db.query(models.User).filter(models.User.id.in_(idlist)).all() if idlist else []
    else:
        us = _scope_users(db, user).all()
    # Anyone checked out today counts as offline, even if a browser tab keeps beating.
    out_today = set()
    if us:
        d = _ist_today()
        for (uid,) in db.query(models.Attendance.user_id).filter(
                models.Attendance.date == d,
                models.Attendance.user_id.in_([u.id for u in us]),
                models.Attendance.check_out_at.isnot(None)).all():
            out_today.add(uid)
    res = {}
    for u in us:
        p = presence.presence_dict(u)
        if u.id in out_today:
            p = {**p, "state": "offline", "idle_seconds": 0}
        res[str(u.id)] = p
    return res


# ================================================================ Excel export

@router.get("/download")
def download_month(month: str | None = None, role: str | None = None, user_id: int | None = None,
                   db: Session = Depends(get_db),
                   user: models.User = Depends(require_roles(*DOWNLOAD_ROLES))):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    data = month_sheet(month, role, user_id, db, user)
    m = data["month"]
    days = data["days"]
    people = data["people"]

    HEAD = PatternFill("solid", fgColor="1D4ED8")
    HEADF = Font(bold=True, color="FFFFFF", size=11)
    THIN = Border(*[Side(style="thin", color="D6DEEA")] * 4)
    COLOR = {"P": "C6EFCE", "L": "FFEB9C", "LV": "BDD7EE", "W": "E2E8F0", "A": "FFC7CE"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Attendance {m}"
    header = ["Emp", "Name", "Role", "Branch"] + [d[8:] for d in days] + \
             ["Present", "Late", "Leave", "Absent", "Week-off", "Work hrs"]
    ws.append(header)
    for c in ws[1]:
        c.fill = HEAD; c.font = HEADF; c.alignment = Alignment(horizontal="center"); c.border = THIN
    for p in people:
        row = [p["emp_code"] or "", p["name"], p["role"], p["branch"] or ""] + \
              [p["days"].get(d, "") for d in days] + \
              [p["present"], p["late"], p["leave"], p["absent"], p["weekoff"], p["worked_hours"]]
        ws.append(row)
        r = ws.max_row
        for i, d in enumerate(days):
            cell = ws.cell(row=r, column=5 + i)
            key = str(cell.value)
            if key in COLOR:
                cell.fill = PatternFill("solid", fgColor=COLOR[key])
            cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "E2"
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=6)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 1, 4), 26)

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    audit.record(db, user, "download", None, entity_type="download",
                 detail=f"Downloaded attendance sheet {m}")
    db.commit()
    fname = f"Attendance_{m}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})
