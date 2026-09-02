"""Daily bank-feedback sheet — maintained per product, one row per case per day.

Callers / back-office give this to the bank daily (or every few days / weekly). Rows
auto-fill from the latest call & field-visit logs and stay editable; manual edits win
and are preserved. Supports per-column filtering, selective row/column download, and
lives updates across viewers.
"""
import io
from datetime import date, datetime, time, timedelta, timezone

from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from .. import models
from ..database import get_db
from ..deps import get_current_user, require_roles
from .cases import _scope
from .realtime import notify_data_changed

router = APIRouter(prefix="/api/feedback", tags=["feedback"])
IST = timezone(timedelta(hours=5, minutes=30))
FEEDBACK_ROLES = ("admin", "manager", "telecaller", "backend", "headoffice", "teamlead")

# ---- The bank's feedback format (order matters — matches the agency's Excel) ----
COLUMNS = [
    {"key": "agency_name", "label": "AGENCY NAME", "type": "text"},
    {"key": "loan_no", "label": "LOAN NO", "type": "text", "readonly": True},
    {"key": "visited", "label": "Visited/Not Visited", "type": "code", "code": "visited"},
    {"key": "dispo_code", "label": "Dispo Code", "type": "code", "code": "dispo_code"},
    {"key": "visit_date", "label": "Visit Date", "type": "date"},
    {"key": "nature_of_business", "label": "Nature Of Business", "type": "code", "code": "nature_of_business"},
    {"key": "default_reason", "label": "Default Reason", "type": "code", "code": "default_reason"},
    {"key": "tc_code", "label": "Tc Code", "type": "code", "code": "tc_code"},
    {"key": "fe_code", "label": "Fe Code", "type": "code", "code": "fe_code"},
    {"key": "tc_final_code", "label": "Tc Final Code", "type": "code", "code": "tc_code"},
    {"key": "fe_final_code", "label": "Fe  Final Code", "type": "code", "code": "fe_code"},
    {"key": "tc_remarks", "label": "Tc Remarks", "type": "text"},
    {"key": "fe_remark", "label": "Fe Remark", "type": "text"},
    {"key": "ptp_date", "label": "PTP Date", "type": "date"},
    # Full merged notes/remarks history (all calls + visits, newest first) — read-only reference.
    {"key": "history", "label": "Remarks History", "type": "text", "readonly": True},
]
# columns actually stored/editable on the FeedbackEntry row (loan_no from the case; history is derived)
STORED = {c["key"] for c in COLUMNS} - {"loan_no", "history"}
# fields kept live from the call / visit logs unless a human has overridden them
LOG_FIELDS = {"visited", "dispo_code", "visit_date", "tc_code", "fe_code", "tc_remarks", "fe_remark", "ptp_date"}

CODES = {
    "visited": ["Visited", "Not Visited"],
    "dispo_code": ["No Contact", "Resolved", "Promise to Pay", "Refuse to Pay", "Call Back / Left Message"],
    "nature_of_business": ["Logistics", "Restaurants", "Gems & Jewellery", "Retail", "Manufacturing",
                           "Healthcare", "IT/Software", "Education", "Finance", "Real Estate",
                           "Automobile", "Agriculture", "Textiles", "Hospitality", "Other"],
    "default_reason": ["Financial Crisis", "Medical Issues", "Business Low", "Temporarily Financial Crisis",
                       "Funds Delay", "Expired", "Multiple Loans", "Bills Pending from Clients", "Business Closed"],
    "tc_code": ["RNR", "CALL BACK", "PTP", "PAID", "RTP", "NO ATTEMPT", "CUSTOMER EXPIRED", "DISPUTE",
                "FRAUD SUSPECT", "PNC", "SNC", "LEFT MESSAGE", "BPTP"],
    "fe_code": ["INCOMPLETE ADDO", "INCOMPLETE ADDR", "CALL BACK", "PTP", "CLAIMS PAID", "SHIFTEDO",
                "SHIFTEDR", "RTP", "SKIP", "NO ATTEMPT", "DOOR LOCK", "CUSTOMER EXPIRED",
                "ENTRY RESTRICT", "FRAUD SUSPECT", "DISPUTE", "BPTP"],
}
_TC = set(CODES["tc_code"])
_FE = set(CODES["fe_code"])

_DISPO_TO_CODE = {          # case/call/visit disposition → bank "Dispo Code"
    "PAID": "Resolved", "RESOLVED": "Resolved",
    "PTP": "Promise to Pay", "BPTP": "Promise to Pay",
    "RTP": "Refuse to Pay", "REFUSED": "Refuse to Pay",
    "RNR": "No Contact", "NC": "No Contact", "NO_CONTACT": "No Contact", "NO CONTACT": "No Contact",
    "CALLBACK": "Call Back / Left Message", "CALL BACK": "Call Back / Left Message",
    "LEFT MESSAGE": "Call Back / Left Message",
}


def _ist_bounds(d: date):
    start = datetime.combine(d, time.min, tzinfo=IST).astimezone(timezone.utc)
    return start, start + timedelta(days=1)


def _fmt_date(dt):
    if not dt:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(IST).strftime("%d-%m-%Y")


def _scope_feedback(db: Session, user, bank, product, branch=None, month_bucket=None):
    q = db.query(models.Case).filter(models.Case.bank == bank, models.Case.product == product,
                                     models.Case.removed.isnot(True))
    if user.role == "manager":
        q = q.filter(models.Case.branch == user.branch)
    elif user.role == "teamlead":
        from .cases import teamlead_case_filter
        q = q.filter(teamlead_case_filter(user))     # a team lead sees their team's feedback
    elif branch:
        q = q.filter(models.Case.branch == branch)
    # Month-wise: keep each month's book separate (this / next / last month).
    if month_bucket in ("current", "next", "last"):
        from .cases import _period_bucket
        q = q.filter(models.Case.period == _period_bucket(month_bucket))
    return q


def _derive(case, call, visit):
    """Log-derived defaults for a case's feedback row (call/visit already filtered to the day)."""
    d = {}
    if visit:
        d["visited"] = "Visited"
        d["visit_date"] = _fmt_date(visit.created_at)
        if visit.note:
            d["fe_remark"] = visit.note
        vd = (visit.disposition or "").upper()
        if vd in _FE:
            d["fe_code"] = vd
    else:
        d["visited"] = "Not Visited"
    if call:
        if call.note:
            d["tc_remarks"] = call.note
        cd = (call.disposition or "").upper()
        if cd in _TC:
            d["tc_code"] = cd
        if call.ptp_date:
            d["ptp_date"] = _fmt_date(call.ptp_date)
    # dispo code from whichever disposition we have
    disp = (visit.disposition if visit else None) or (call.disposition if call else None) or case.disposition
    if disp:
        code = _DISPO_TO_CODE.get(str(disp).upper())
        if code:
            d["dispo_code"] = code
    return d


def _latest_by_case(rows):
    out = {}
    for r in rows:                       # rows already ordered oldest→newest, keep last
        out[r.case_id] = r
    return out


def _row(entry, case):
    r = {"id": entry.id, "case_id": case.id, "loan_no": case.account_no,
         "customer": case.customer_name, "phone": case.phone}
    for k in STORED:
        r[k] = getattr(entry, k)
    return r


@router.get("/config")
def feedback_config(user: models.User = Depends(require_roles(*FEEDBACK_ROLES))):
    return {"columns": COLUMNS, "codes": CODES}


@router.get("")
def get_feedback(bank: str, product: str, day: str | None = None, branch: str | None = None,
                 month_bucket: str | None = None,
                 db: Session = Depends(get_db),
                 user: models.User = Depends(require_roles(*FEEDBACK_ROLES))):
    try:
        d = datetime.strptime(day, "%Y-%m-%d").date() if day else datetime.now(IST).date()
    except ValueError:
        d = datetime.now(IST).date()

    cases = _scope_feedback(db, user, bank, product, branch, month_bucket).order_by(models.Case.customer_name).all()
    ids = [c.id for c in cases]
    if not ids:
        return {"day": d.isoformat(), "bank": bank, "product": product, "count": 0, "rows": []}

    start, end = _ist_bounds(d)
    calls = _latest_by_case(db.query(models.CallLog).filter(
        models.CallLog.case_id.in_(ids), models.CallLog.created_at >= start,
        models.CallLog.created_at < end).order_by(models.CallLog.created_at.asc()).all())
    visits = _latest_by_case(db.query(models.Visit).filter(
        models.Visit.case_id.in_(ids), models.Visit.created_at >= start,
        models.Visit.created_at < end).order_by(models.Visit.created_at.asc()).all())
    existing = {e.case_id: e for e in db.query(models.FeedbackEntry).filter(
        models.FeedbackEntry.case_id.in_(ids), models.FeedbackEntry.day == d).all()}

    rows, dirty = [], False
    for c in cases:
        derived = _derive(c, calls.get(c.id), visits.get(c.id))
        e = existing.get(c.id)
        if not e:
            e = models.FeedbackEntry(case_id=c.id, bank=bank, product=product, day=d,
                                     agency_name=(c.branch or None), edited={}, **derived)
            db.add(e); dirty = True
        else:
            ed = e.edited or {}
            for f in LOG_FIELDS:                       # refresh live fields the caller hasn't touched
                if not ed.get(f) and derived.get(f) and getattr(e, f) != derived[f]:
                    setattr(e, f, derived[f]); dirty = True
        rows.append((e, c))
    if dirty:
        db.commit()
    # Full merged notes/remarks history (all time, newest first) per case for the History column.
    from ..notes import case_notes_map, join_notes
    nmap = case_notes_map(db, ids)
    out_rows = []
    for e, c in rows:
        r = _row(e, c)
        r["history"] = join_notes(nmap.get(c.id, []))
        out_rows.append(r)
    return {"day": d.isoformat(), "bank": bank, "product": product,
            "count": len(rows), "rows": out_rows}


@router.patch("/{entry_id}")
def edit_feedback(entry_id: int, field: str = Body(...), value: str | None = Body(None),
                  db: Session = Depends(get_db),
                  user: models.User = Depends(require_roles(*FEEDBACK_ROLES))):
    if field not in STORED:
        raise HTTPException(status_code=400, detail=f"'{field}' is not editable")
    e = db.query(models.FeedbackEntry).filter(models.FeedbackEntry.id == entry_id).first()
    if not e:
        raise HTTPException(status_code=404, detail="Row not found")
    if user.role == "manager":
        case = db.query(models.Case).filter(models.Case.id == e.case_id).first()
        if case and case.branch != user.branch:
            raise HTTPException(status_code=403, detail="Not in your branch")
    setattr(e, field, (value or None))
    ed = dict(e.edited or {})
    ed[field] = True                                   # remember the human override
    e.edited = ed
    e.updated_by = user.id
    db.commit()
    notify_data_changed(e.bank, e.product)
    return {"ok": True}


@router.post("/refresh")
def refresh_feedback(bank: str, product: str, day: str | None = None, branch: str | None = None,
                     month_bucket: str | None = None,
                     db: Session = Depends(get_db),
                     user: models.User = Depends(require_roles(*FEEDBACK_ROLES))):
    """Force-pull the log-derived fields from the latest call/visit, overriding prior auto values
    (manual edits are still kept)."""
    try:
        d = datetime.strptime(day, "%Y-%m-%d").date() if day else datetime.now(IST).date()
    except ValueError:
        d = datetime.now(IST).date()
    cases = _scope_feedback(db, user, bank, product, branch, month_bucket).all()
    ids = [c.id for c in cases]
    if not ids:
        return {"ok": True, "updated": 0}
    start, end = _ist_bounds(d)
    calls = _latest_by_case(db.query(models.CallLog).filter(
        models.CallLog.case_id.in_(ids), models.CallLog.created_at >= start,
        models.CallLog.created_at < end).order_by(models.CallLog.created_at.asc()).all())
    visits = _latest_by_case(db.query(models.Visit).filter(
        models.Visit.case_id.in_(ids), models.Visit.created_at >= start,
        models.Visit.created_at < end).order_by(models.Visit.created_at.asc()).all())
    existing = {e.case_id: e for e in db.query(models.FeedbackEntry).filter(
        models.FeedbackEntry.case_id.in_(ids), models.FeedbackEntry.day == d).all()}
    n = 0
    for c in cases:
        derived = _derive(c, calls.get(c.id), visits.get(c.id))
        e = existing.get(c.id)
        if not e:
            e = models.FeedbackEntry(case_id=c.id, bank=bank, product=product, day=d,
                                     agency_name=(c.branch or None), edited={}, **derived)
            db.add(e); n += 1; continue
        ed = e.edited or {}
        for f in LOG_FIELDS:
            if not ed.get(f) and derived.get(f):
                setattr(e, f, derived[f]); n += 1
    db.commit()
    notify_data_changed(bank, product)
    return {"ok": True, "updated": n}


@router.get("/download")
def download_feedback(bank: str, product: str, day: str | None = None, branch: str | None = None,
                      month_bucket: str | None = None,
                      columns: str | None = None, ids: str | None = None,
                      db: Session = Depends(get_db),
                      user: models.User = Depends(require_roles(*FEEDBACK_ROLES))):
    """Excel export. `columns` = comma-separated column keys (default all); `ids` = comma-separated
    entry ids to include only selected rows."""
    from openpyxl import Workbook
    try:
        d = datetime.strptime(day, "%Y-%m-%d").date() if day else datetime.now(IST).date()
    except ValueError:
        d = datetime.now(IST).date()

    payload = get_feedback(bank, product, d.isoformat(), branch, month_bucket, db, user)
    rows = payload["rows"]
    if ids:
        keep = {int(x) for x in ids.split(",") if x.strip().isdigit()}
        rows = [r for r in rows if r["id"] in keep]
    keys = [k for k in columns.split(",")] if columns else [c["key"] for c in COLUMNS]
    cols = [c for c in COLUMNS if c["key"] in keys] or COLUMNS

    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"
    ws.append([c["label"] for c in cols])
    for r in rows:
        ws.append([r.get(c["key"], "") if r.get(c["key"]) is not None else "" for c in cols])
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 12), 42)

    # a "Codes" sheet so the bank's file keeps the same dropdown lists
    cs = wb.create_sheet("Codes")
    code_cols = [(c["label"], CODES[c["code"]]) for c in cols if c.get("code")]
    cs.append([lbl for lbl, _ in code_cols])
    for i in range(max((len(v) for _, v in code_cols), default=0)):
        cs.append([(v[i] if i < len(v) else "") for _, v in code_cols])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fname = f"feedback_{bank}_{product}_{d.isoformat()}.xlsx".replace(" ", "_")
    from .. import audit as _audit
    _audit.record(db, user, "download", None, entity_type="download",
                  detail=f"Downloaded Bank Feedback — {bank}/{product} {d.isoformat()}")
    db.commit()
    return StreamingResponse(
        out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})
