import io
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session
from openpyxl import Workbook

from .. import models, schemas
from ..database import get_db
from ..deps import get_current_user, require_roles

router = APIRouter(prefix="/api/tracking", tags=["tracking"])

# Roles allowed to view a field officer's route history / map.
HISTORY_ROLES = ("admin", "manager", "headoffice", "teamlead")


def _authorize_officer_view(db: Session, viewer: models.User, officer_id: int):
    """Managers may only view officers in their branch; team leads only their own
    team members. Admin, head office & tech-support see everyone. Raises 403 otherwise."""
    if viewer.role in ("admin", "headoffice", "techsupport"):
        return
    if viewer.role == "manager":
        off = db.query(models.User).filter(models.User.id == officer_id).first()
        if not off or off.branch != viewer.branch:
            raise HTTPException(status_code=403, detail="Not in your branch")
        return
    if viewer.role == "teamlead":
        from .cases import _scope_user_ids
        if officer_id not in set(_scope_user_ids(db, viewer)):
            raise HTTPException(status_code=403, detail="Not one of your team members")
        return
    raise HTTPException(status_code=403, detail="Not allowed")


# Server-side ping coalescer. Field apps send a keep-alive ping every few seconds even while the
# officer stands still, which floods the DB's single SQLite writer and slows other writes (e.g. a
# visit submit) to a crawl. We only need to STORE a new point when the officer has actually moved,
# or once every _PING_MIN_GAP_S so presence stays fresh. When a ping is "stationary + too soon" we
# skip the DB write entirely and just echo it back. Movement (>= _PING_MIN_MOVE_M) is always stored,
# so route history keeps full detail. In-memory only (best-effort); harmless on Postgres too.
_PING_LAST: dict[int, tuple] = {}          # officer_id -> (epoch_seconds, lat, lng)
_PING_MIN_GAP_S = 12.0                       # while stationary, store at most one ping per 12s
_PING_MIN_MOVE_M = 25.0                      # always store if moved at least 25 m


def _ping_moved_m(a_lat, a_lng, b_lat, b_lng) -> float:
    import math
    r = 6371000.0
    p1, p2 = math.radians(a_lat), math.radians(b_lat)
    dphi = math.radians(b_lat - a_lat)
    dl = math.radians(b_lng - a_lng)
    h = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(h))


@router.post("/ping", response_model=schemas.PingOut)
def ping(body: schemas.PingCreate, db: Session = Depends(get_db),
         user: models.User = Depends(require_roles("fos", "admin"))):
    import time as _time
    now = _time.time()
    prev = _PING_LAST.get(user.id)
    if (prev and body.latitude is not None and body.longitude is not None
            and (now - prev[0]) < _PING_MIN_GAP_S
            and _ping_moved_m(prev[1], prev[2], body.latitude, body.longitude) < _PING_MIN_MOVE_M):
        # Stationary and pinged again too soon — skip the write, just acknowledge. (Don't refresh
        # the timestamp, so a real write still lands once _PING_MIN_GAP_S has elapsed.)
        return {
            "id": 0, "officer_id": user.id, "latitude": body.latitude, "longitude": body.longitude,
            "accuracy": body.accuracy, "active_case_id": body.active_case_id,
            "created_at": datetime.now(timezone.utc),
        }
    p = models.LocationPing(
        officer_id=user.id, latitude=body.latitude, longitude=body.longitude,
        accuracy=body.accuracy, speed=body.speed, active_case_id=body.active_case_id,
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    if body.latitude is not None and body.longitude is not None:
        _PING_LAST[user.id] = (now, body.latitude, body.longitude)
    return p


def _extract_coords(obj):
    """Pull {latitude, longitude, accuracy, speed} out of a Transistor location object,
    which may be shaped as {coords:{...}} or nested under {location:{coords:{...}}}."""
    if not isinstance(obj, dict):
        return None
    c = obj.get("coords") or (obj.get("location") or {}).get("coords") or obj
    lat, lng = c.get("latitude"), c.get("longitude")
    if lat is None or lng is None:
        return None
    return dict(latitude=float(lat), longitude=float(lng),
                accuracy=c.get("accuracy"), speed=c.get("speed"))


@router.post("/ping-native")
async def ping_native(request: Request, db: Session = Depends(get_db),
                      user: models.User = Depends(require_roles("fos", "admin"))):
    """Receives locations POSTed directly by the native Android background-geolocation
    service (Transistor auto-sync). Accepts a single location or a batch, in the plugin's
    default shape. Stores each as a LocationPing for the authenticated officer."""
    try:
        data = await request.json()
    except Exception:
        data = None
    items = []
    if isinstance(data, list):
        items = data
    elif isinstance(data, dict):
        loc = data.get("location", data)
        items = loc if isinstance(loc, list) else [loc]
    stored = 0
    for it in items:
        c = _extract_coords(it)
        if not c:
            continue
        db.add(models.LocationPing(officer_id=user.id, **c))
        stored += 1
    db.commit()
    return {"stored": stored}


@router.get("/live", response_model=list[schemas.OfficerLocation])
def live(minutes: int = 30, db: Session = Depends(get_db),
         viewer: models.User = Depends(require_roles("admin", "manager", "headoffice"))):
    """Latest known position of every field officer seen in the last N minutes.
    A branch manager only sees officers in their own branch."""
    since = datetime.now(timezone.utc) - timedelta(minutes=minutes)
    sub = (
        db.query(models.LocationPing.officer_id,
                 func.max(models.LocationPing.created_at).label("mx"))
        .filter(models.LocationPing.created_at >= since)
        .group_by(models.LocationPing.officer_id)
        .subquery()
    )
    q = (
        db.query(models.LocationPing, models.User.name, models.User.branch, models.User.banks)
        .join(sub, (models.LocationPing.officer_id == sub.c.officer_id) &
              (models.LocationPing.created_at == sub.c.mx))
        .join(models.User, models.User.id == models.LocationPing.officer_id)
    )
    if viewer.role == "manager":
        q = q.filter(models.User.branch == viewer.branch)
    rows = q.all()
    return [
        schemas.OfficerLocation(
            officer_id=p.officer_id, name=name, latitude=p.latitude, longitude=p.longitude,
            accuracy=p.accuracy, active_case_id=p.active_case_id, last_seen=p.created_at,
            branch=branch, banks=banks or [],
        )
        for p, name, branch, banks in rows
    ]


@router.get("/officer/{officer_id}/trail", response_model=list[schemas.PingOut])
def trail(officer_id: int, minutes: int = 240, db: Session = Depends(get_db),
          viewer: models.User = Depends(require_roles(*HISTORY_ROLES))):
    _authorize_officer_view(db, viewer, officer_id)
    since = datetime.now(timezone.utc) - timedelta(minutes=minutes)
    return (
        db.query(models.LocationPing)
        .filter(models.LocationPing.officer_id == officer_id,
                models.LocationPing.created_at >= since)
        .order_by(models.LocationPing.created_at.asc())
        .all()
    )


# ---------------- 90-day (3-month) route history ----------------
IST = timezone(timedelta(hours=5, minutes=30))   # India Standard Time


def _ist_today():
    return datetime.now(IST).date()
RETENTION_DAYS = 90


def _as_utc(dt: datetime) -> datetime:
    """Normalise a possibly-naive stored datetime to aware UTC."""
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def prune_old_pings(db: Session, days: int = RETENTION_DAYS) -> int:
    cutoff = datetime.now(timezone.utc) - timedelta(days=days)
    deleted = (
        db.query(models.LocationPing)
        .filter(models.LocationPing.created_at < cutoff)
        .delete(synchronize_session=False)
    )
    db.commit()
    return deleted


def _haversine_km(a, b):
    import math
    r = 6371.0
    p1, p2 = math.radians(a[0]), math.radians(b[0])
    dphi = math.radians(b[0] - a[0])
    dl = math.radians(b[1] - a[1])
    h = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(h))


def clean_route(pings, max_accuracy_m=80.0, max_speed_kmh=140.0):
    """Drop GPS junk so a stationary/slow officer doesn't draw spikes across the map:
      * fixes with poor reported accuracy (drift, indoor, cold start), and
      * 'teleport' outliers — a big jump in a tiny time gap (impossible speed).
    Keeps chronological order; compares each candidate to the last KEPT good point."""
    kept = []
    for p in pings:
        if p.latitude is None or p.longitude is None:
            continue
        if p.accuracy is not None and p.accuracy > max_accuracy_m:
            continue                                  # low-accuracy fix
        if kept:
            prev = kept[-1]
            dist_km = _haversine_km((prev.latitude, prev.longitude), (p.latitude, p.longitude))
            dt = 0.0
            if p.created_at and prev.created_at:
                dt = (_as_utc(p.created_at) - _as_utc(prev.created_at)).total_seconds()
            if dt > 0 and dist_km > 0.15 and (dist_km / (dt / 3600.0)) > max_speed_kmh:
                continue                              # impossible jump → bad fix
        kept.append(p)
    return kept


@router.get("/officer/{officer_id}/history-dates")
def history_dates(officer_id: int, days: int = RETENTION_DAYS,
                  db: Session = Depends(get_db),
                  viewer: models.User = Depends(require_roles(*HISTORY_ROLES))):
    """Days (IST) in the last `days` (default 90 = 3 months) on which this officer has recorded a route."""
    _authorize_officer_view(db, viewer, officer_id)
    prune_old_pings(db)  # keep only the last 90 days (3 months)
    since = datetime.now(timezone.utc) - timedelta(days=days)
    pings = (
        db.query(models.LocationPing)
        .filter(models.LocationPing.officer_id == officer_id,
                models.LocationPing.created_at >= since)
        .order_by(models.LocationPing.created_at.asc())
        .all()
    )
    pings = clean_route(pings)             # de-jitter so per-day distance isn't inflated
    by_day = {}
    for p in pings:
        local = _as_utc(p.created_at).astimezone(IST)
        key = local.strftime("%Y-%m-%d")
        d = by_day.setdefault(key, {"date": key, "points": 0, "first": None, "last": None, "coords": []})
        d["points"] += 1
        d["last"] = local
        if d["first"] is None:
            d["first"] = local
        d["coords"].append((p.latitude, p.longitude))
    out = []
    for key in sorted(by_day, reverse=True):
        d = by_day[key]
        dist = 0.0
        for i in range(1, len(d["coords"])):
            dist += _haversine_km(d["coords"][i - 1], d["coords"][i])
        out.append({
            "date": d["date"],
            "points": d["points"],
            "distance_km": round(dist, 2),
            "first_seen": d["first"].strftime("%H:%M"),
            "last_seen": d["last"].strftime("%H:%M"),
        })
    return out


@router.get("/officer/{officer_id}/route", response_model=list[schemas.PingOut])
def route_for_date(officer_id: int, date: str, db: Session = Depends(get_db),
                   viewer: models.User = Depends(require_roles(*HISTORY_ROLES))):
    """Full ordered route for one IST calendar day (date = 'YYYY-MM-DD')."""
    _authorize_officer_view(db, viewer, officer_id)
    try:
        y, m, d = (int(x) for x in date.split("-"))
        start_ist = datetime(y, m, d, 0, 0, tzinfo=IST)
    except Exception:
        from fastapi import HTTPException
        raise HTTPException(status_code=400, detail="date must be YYYY-MM-DD")
    start_utc = start_ist.astimezone(timezone.utc)
    end_utc = start_utc + timedelta(days=1)
    pings = (
        db.query(models.LocationPing)
        .filter(models.LocationPing.officer_id == officer_id,
                models.LocationPing.created_at >= start_utc,
                models.LocationPing.created_at < end_utc)
        .order_by(models.LocationPing.created_at.asc())
        .all()
    )
    return clean_route(pings)             # strip GPS jitter so the drawn route isn't spiky


@router.get("/roster")
def fos_roster(date: str | None = None, db: Session = Depends(get_db),
               user: models.User = Depends(require_roles("admin", "manager", "headoffice", "teamlead"))):
    """All field officers split into ACTIVE (shared location on the day) vs INACTIVE (no
    tracking that day), for today or any past date (date='YYYY-MM-DD'). Powers the live-map popup."""
    if date:
        try:
            y, m, d = (int(x) for x in date.split("-"))
            day = datetime(y, m, d, tzinfo=IST).date()
        except Exception:
            from fastapi import HTTPException
            raise HTTPException(status_code=400, detail="date must be YYYY-MM-DD")
    else:
        day = _ist_today()
    start_utc = datetime(day.year, day.month, day.day, tzinfo=IST).astimezone(timezone.utc)
    end_utc = start_utc + timedelta(days=1)

    q = db.query(models.User).filter(models.User.role == "fos", models.User.is_active.isnot(False))
    if user.role == "manager" and user.branch:
        q = q.filter(models.User.branch == user.branch)
    elif user.role == "teamlead":
        q = q.filter(models.User.team_lead_id == user.id)
    foses = q.order_by(models.User.name).all()
    ids = [u.id for u in foses]

    pings = []
    if ids:
        pings = (db.query(models.LocationPing)
                 .filter(models.LocationPing.officer_id.in_(ids),
                         models.LocationPing.created_at >= start_utc,
                         models.LocationPing.created_at < end_utc)
                 .order_by(models.LocationPing.created_at.asc()).all())
    by_officer: dict = {}
    for p in pings:
        by_officer.setdefault(p.officer_id, []).append(p)

    def _info(u):
        base = {"id": u.id, "name": u.name, "emp_code": u.emp_code, "branch": u.branch,
                "location": u.location, "phone": u.phone, "active": u.id in by_officer}
        ps = by_officer.get(u.id)
        if ps:
            cleaned = clean_route(ps) or ps
            dist = 0.0
            for a, b in zip(cleaned, cleaned[1:]):
                dist += _haversine_km((a.latitude, a.longitude), (b.latitude, b.longitude))
            base.update({
                "pings": len(ps),
                "first_seen": _as_utc(ps[0].created_at).isoformat(),
                "last_seen": _as_utc(ps[-1].created_at).isoformat(),
                "distance_km": round(dist, 2),
            })
        return base

    active = [_info(u) for u in foses if u.id in by_officer]
    inactive = [_info(u) for u in foses if u.id not in by_officer]
    return {
        "date": day.isoformat(),
        "is_today": day == _ist_today(),
        "total": len(foses), "active_count": len(active), "inactive_count": len(inactive),
        "active": active, "inactive": inactive,
    }


@router.get("/distance-report")
def distance_report(start: str, end: str, officer_id: int | None = None,
                    db: Session = Depends(get_db),
                    admin: models.User = Depends(require_roles("admin", "manager"))):
    """Excel attendance/distance report: per-officer per-day km travelled and active hours
    across an inclusive IST date range (start & end = 'YYYY-MM-DD')."""
    try:
        sy, sm, sd = (int(x) for x in start.split("-"))
        ey, em, ed = (int(x) for x in end.split("-"))
        start_utc = datetime(sy, sm, sd, 0, 0, tzinfo=IST).astimezone(timezone.utc)
        end_utc = (datetime(ey, em, ed, 0, 0, tzinfo=IST) + timedelta(days=1)).astimezone(timezone.utc)
    except Exception:
        raise HTTPException(status_code=400, detail="start and end must be YYYY-MM-DD")
    if end_utc <= start_utc:
        raise HTTPException(status_code=400, detail="end date must be on or after start date")

    q = db.query(models.LocationPing).filter(
        models.LocationPing.created_at >= start_utc, models.LocationPing.created_at < end_utc)
    if officer_id:
        q = q.filter(models.LocationPing.officer_id == officer_id)
    pings = q.order_by(models.LocationPing.officer_id, models.LocationPing.created_at.asc()).all()

    names = {u.id: u.name for u in db.query(models.User).all()}
    groups = {}
    for p in pings:
        local = _as_utc(p.created_at).astimezone(IST)
        key = (p.officer_id, local.strftime("%Y-%m-%d"))
        g = groups.setdefault(key, {"coords": [], "first": local, "last": local})
        g["coords"].append((p.latitude, p.longitude))
        if local < g["first"]:
            g["first"] = local
        if local > g["last"]:
            g["last"] = local

    daily, per_officer = [], {}
    for (oid, date), g in sorted(groups.items(), key=lambda kv: (names.get(kv[0][0], ""), kv[0][1])):
        dist = sum(_haversine_km(g["coords"][i - 1], g["coords"][i]) for i in range(1, len(g["coords"])))
        active_h = round((g["last"] - g["first"]).total_seconds() / 3600, 2)
        daily.append([names.get(oid, str(oid)), date, len(g["coords"]), round(dist, 2),
                      g["first"].strftime("%H:%M"), g["last"].strftime("%H:%M"), active_h])
        po = per_officer.setdefault(oid, {"days": 0, "dist": 0.0, "active": 0.0})
        po["days"] += 1; po["dist"] += dist; po["active"] += active_h

    wb = Workbook()
    ws = wb.active; ws.title = "Daily"
    ws.append(["Officer", "Date", "GPS Points", "Distance (km)", "First Seen", "Last Seen", "Active Hours"])
    for r in daily:
        ws.append(r)
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Officer", "Active Days", "Total Distance (km)", "Total Active Hours", "Avg km/day"])
    for oid, po in sorted(per_officer.items(), key=lambda kv: names.get(kv[0], "")):
        ws2.append([names.get(oid, str(oid)), po["days"], round(po["dist"], 2), round(po["active"], 2),
                    round(po["dist"] / po["days"], 2) if po["days"] else 0])
    for sheet in (ws, ws2):
        for cell in sheet[1]:
            cell.font = cell.font.copy(bold=True)
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 40)

    out = io.BytesIO(); wb.save(out)
    fname = f"SSD_Attendance_{start}_to_{end}.xlsx"
    from .. import audit as _audit
    _audit.record(db, admin, "download", None, entity_type="download",
                  detail=f"Downloaded attendance / distance report {start} → {end}")
    db.commit()
    return StreamingResponse(
        io.BytesIO(out.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


@router.get("/officer/{officer_id}/route.csv")
def route_csv(officer_id: int, date: str, db: Session = Depends(get_db),
              viewer: models.User = Depends(require_roles(*HISTORY_ROLES))):
    """Download a field officer's full GPS route for one IST day as CSV
    (sequence, timestamp, latitude, longitude, accuracy, speed)."""
    _authorize_officer_view(db, viewer, officer_id)
    try:
        y, m, d = (int(x) for x in date.split("-"))
        start_utc = datetime(y, m, d, 0, 0, tzinfo=IST).astimezone(timezone.utc)
    except Exception:
        raise HTTPException(status_code=400, detail="date must be YYYY-MM-DD")
    end_utc = start_utc + timedelta(days=1)
    pings = (db.query(models.LocationPing)
             .filter(models.LocationPing.officer_id == officer_id,
                     models.LocationPing.created_at >= start_utc,
                     models.LocationPing.created_at < end_utc)
             .order_by(models.LocationPing.created_at.asc()).all())
    officer = db.query(models.User).filter(models.User.id == officer_id).first()
    name = officer.name if officer else str(officer_id)

    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["officer", "date", "seq", "timestamp_ist", "latitude", "longitude", "accuracy_m", "speed"])
    for i, p in enumerate(pings, start=1):
        ts = _as_utc(p.created_at).astimezone(IST).strftime("%Y-%m-%d %H:%M:%S")
        w.writerow([name, date, i, ts, p.latitude, p.longitude, p.accuracy or "", p.speed or ""])

    safe = "".join(ch for ch in name if ch.isalnum() or ch in "-_") or str(officer_id)
    from .. import audit as _audit
    _audit.record(db, viewer, "download", None, entity_type="download", target_user_id=officer_id,
                  detail=f"Downloaded route CSV for {name} on {date}")
    db.commit()
    return StreamingResponse(
        io.BytesIO(buf.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=route_{safe}_{date}.csv"},
    )


@router.get("/available-dates")
def available_dates(officer_ids: str = "all", branch: str | None = None,
                    days: int = RETENTION_DAYS, db: Session = Depends(get_db),
                    viewer: models.User = Depends(require_roles(*HISTORY_ROLES))):
    """Union of IST calendar days (last `days`) on which any of the given officers recorded a
    route — powers the multi-date picker in the bulk export panel."""
    fq = db.query(models.User.id).filter(models.User.role == "fos")
    if viewer.role == "manager" and viewer.branch:
        fq = fq.filter(models.User.branch == viewer.branch)
    elif viewer.role == "teamlead":
        from .cases import _scope_user_ids
        fq = fq.filter(models.User.id.in_(list(_scope_user_ids(db, viewer)) or [-1]))
    if branch:
        fq = fq.filter(models.User.branch == branch)
    if officer_ids and officer_ids != "all":
        want = {int(x) for x in officer_ids.split(",") if x.strip().isdigit()}
        fq = fq.filter(models.User.id.in_(want or {-1}))
    oids = [r[0] for r in fq.all()]
    if not oids:
        return []
    since = datetime.now(timezone.utc) - timedelta(days=days)
    rows = (db.query(models.LocationPing.created_at)
            .filter(models.LocationPing.officer_id.in_(oids),
                    models.LocationPing.created_at >= since).all())
    return sorted({_as_utc(r[0]).astimezone(IST).strftime("%Y-%m-%d") for r in rows}, reverse=True)


@router.get("/export")
def export_routes(officer_ids: str = "all", dates: str = "all", branch: str | None = None,
                  days: int = RETENTION_DAYS, db: Session = Depends(get_db),
                  viewer: models.User = Depends(require_roles(*HISTORY_ROLES))):
    """Multi-officer, multi-date route + visit export as ONE Excel workbook.

    `officer_ids` = comma-separated user ids, or 'all'. `dates` = comma-separated YYYY-MM-DD,
    or 'all' (every day in the last `days` on which any selected officer recorded a route).
    Layout: one worksheet per date; within a sheet each officer is a stacked block — a name +
    route-summary header row, then that officer's field-visit table.
    """
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    # ---- resolve the officer set (scoped to what the viewer may see) ----
    fq = db.query(models.User).filter(models.User.role == "fos")
    if viewer.role == "manager" and viewer.branch:
        fq = fq.filter(models.User.branch == viewer.branch)
    elif viewer.role == "teamlead":
        from .cases import _scope_user_ids
        fq = fq.filter(models.User.id.in_(list(_scope_user_ids(db, viewer)) or [-1]))
    if branch:
        fq = fq.filter(models.User.branch == branch)
    if officer_ids and officer_ids != "all":
        want = {int(x) for x in officer_ids.split(",") if x.strip().isdigit()}
        fq = fq.filter(models.User.id.in_(want or {-1}))
    officers = fq.order_by(models.User.name).all()
    if not officers:
        raise HTTPException(status_code=404, detail="No field officers match the selection")
    for o in officers:                                  # enforce per-officer visibility
        _authorize_officer_view(db, viewer, o.id)
    oids = [o.id for o in officers]

    # ---- resolve the date set ----
    since = datetime.now(timezone.utc) - timedelta(days=days)
    if dates and dates != "all":
        date_list = sorted({x.strip() for x in dates.split(",") if x.strip()})
    else:
        pings = (db.query(models.LocationPing.created_at)
                 .filter(models.LocationPing.officer_id.in_(oids),
                         models.LocationPing.created_at >= since).all())
        date_list = sorted({_as_utc(p[0]).astimezone(IST).strftime("%Y-%m-%d") for p in pings})
    if not date_list:
        raise HTTPException(status_code=404, detail="No route data for the selection")

    def _day_bounds(dstr):
        y, m, d = (int(x) for x in dstr.split("-"))
        s = datetime(y, m, d, 0, 0, tzinfo=IST).astimezone(timezone.utc)
        return s, s + timedelta(days=1)

    def as_ist_hm(dt):
        if dt and dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(IST).strftime("%H:%M") if dt else ""

    VIS_COLS = ["#", "Time", "Customer", "Account", "Bank", "Product", "Phone",
                "Disposition", "Paid", "Amount", "N/S", "Off-loc", "Distance (m)", "Note"]
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    name_fill = PatternFill("solid", fgColor="DDEBF7")
    col_fill = PatternFill("solid", fgColor="F2F2F2")
    white_bold = Font(bold=True, color="FFFFFF")
    name_font = Font(bold=True, size=12, color="1F4E79")

    wb = Workbook()
    wb.remove(wb.active)
    total_visits = 0
    for dstr in date_list:
        start_utc, end_utc = _day_bounds(dstr)
        ws = wb.create_sheet(title=dstr[:31])
        row = 1
        # per-day pings & visits for all officers in two queries
        day_pings = (db.query(models.LocationPing)
                     .filter(models.LocationPing.officer_id.in_(oids),
                             models.LocationPing.created_at >= start_utc,
                             models.LocationPing.created_at < end_utc)
                     .order_by(models.LocationPing.created_at.asc()).all())
        pings_by_off = {}
        for p in day_pings:
            pings_by_off.setdefault(p.officer_id, []).append(p)
        day_visits = (db.query(models.Visit)
                      .filter(models.Visit.officer_id.in_(oids),
                              models.Visit.created_at >= start_utc,
                              models.Visit.created_at < end_utc)
                      .order_by(models.Visit.created_at.asc()).all())
        vis_by_off = {}
        for v in day_visits:
            vis_by_off.setdefault(v.officer_id, []).append(v)
        cids = list({v.case_id for v in day_visits})
        cases = {c.id: c for c in db.query(models.Case).filter(models.Case.id.in_(cids)).all()} if cids else {}

        for o in officers:
            ps = pings_by_off.get(o.id, [])
            vs = vis_by_off.get(o.id, [])
            if not ps and not vs:
                continue                                # nothing this day for this officer — skip
            cleaned = clean_route(ps) or ps
            dist = 0.0
            for a, b in zip(cleaned, cleaned[1:]):
                dist += _haversine_km((a.latitude, a.longitude), (b.latitude, b.longitude))
            first = as_ist_hm(ps[0].created_at) if ps else "—"
            last = as_ist_hm(ps[-1].created_at) if ps else "—"
            collected = round(sum(float(v.amount_collected or 0) for v in vs), 2)

            # ---- officer name header ----
            c0 = ws.cell(row=row, column=1,
                         value=f"{o.name}"
                               + (f"  ·  {o.emp_code}" if o.emp_code else "")
                               + (f"  ·  {o.branch}" if o.branch else ""))
            c0.font = name_font
            for cc in range(1, len(VIS_COLS) + 1):
                ws.cell(row=row, column=cc).fill = name_fill
            row += 1
            # ---- route summary line ----
            ws.cell(row=row, column=1,
                    value=f"Route: {dist:.2f} km · {len(ps)} points · {first}–{last}   |   "
                          f"Visits: {len(vs)} · Collected ₹{collected:.0f}").font = Font(italic=True, size=10)
            row += 1
            # ---- visits table ----
            for i, col in enumerate(VIS_COLS, start=1):
                cell = ws.cell(row=row, column=i, value=col)
                cell.font = white_bold
                cell.fill = hdr_fill
                cell.alignment = Alignment(horizontal="center")
            row += 1
            if vs:
                for idx, v in enumerate(vs, start=1):
                    c = cases.get(v.case_id)
                    off_loc = (v.distance_from_case_m is not None and v.distance_from_case_m > 300)
                    vals = [idx, as_ist_hm(v.created_at),
                            c.customer_name if c else "", c.account_no if c else "",
                            c.bank if c else "", c.product if c else "", c.phone if c else "",
                            v.disposition or "", "Yes" if v.paid else "",
                            float(v.amount_collected or 0), (c.norm_stab if c else "") or "",
                            "Yes" if off_loc else "", v.distance_from_case_m or "", v.note or ""]
                    for i, val in enumerate(vals, start=1):
                        ws.cell(row=row, column=i, value=val)
                    row += 1
                    total_visits += 1
            else:
                ws.cell(row=row, column=1, value="No field visits logged this day.").font = Font(italic=True, color="808080")
                row += 1
            row += 1                                     # blank spacer between officers

        if row == 1:                                     # nobody had data this day
            ws.cell(row=1, column=1, value="No route or visit data for any selected officer on this date.")
        # column widths
        widths = [5, 7, 22, 16, 12, 14, 14, 14, 6, 10, 6, 8, 12, 40]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A1"

    if not wb.sheetnames:
        ws = wb.create_sheet("Export")
        ws.cell(row=1, column=1, value="No data for the selected officers and dates.")

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    n_off, n_day = len(officers), len(date_list)
    fname = (f"route_export_{n_off}fos_{n_day}days_{date_list[0]}_to_{date_list[-1]}.xlsx"
             if n_day > 1 else f"route_export_{n_off}fos_{date_list[0]}.xlsx")
    from .. import audit as _audit
    _audit.record(db, viewer, "download", None, entity_type="download",
                  detail=f"Bulk route export — {n_off} officer(s) × {n_day} day(s) "
                         f"({date_list[0]}…{date_list[-1]}), {total_visits} visit rows")
    db.commit()
    return StreamingResponse(
        out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/me/today")
def my_today(db: Session = Depends(get_db),
             user: models.User = Depends(require_roles("fos", "admin"))):
    """The logged-in field officer's own GPS route for today (IST) — points + distance,
    for the FO's live map."""
    today = _ist_today()
    start_utc = datetime(today.year, today.month, today.day, 0, 0, tzinfo=IST).astimezone(timezone.utc)
    end_utc = start_utc + timedelta(days=1)
    pings = (db.query(models.LocationPing)
             .filter(models.LocationPing.officer_id == user.id,
                     models.LocationPing.created_at >= start_utc,
                     models.LocationPing.created_at < end_utc)
             .order_by(models.LocationPing.created_at.asc()).all())
    pings = clean_route(pings)                         # strip GPS jitter/outliers
    pts = [{"lat": p.latitude, "lng": p.longitude,
            "at": _as_utc(p.created_at).astimezone(IST).strftime("%H:%M")} for p in pings]
    coords = [(p.latitude, p.longitude) for p in pings]
    dist = sum(_haversine_km(coords[i - 1], coords[i]) for i in range(1, len(coords)))
    return {"count": len(pts), "distance_km": round(dist, 2), "points": pts}
