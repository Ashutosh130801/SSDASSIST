"""Manpower / HR directory.

- Admin, Head Office and HR can browse the full employee directory (role + location
  filters) and download it as Excel.
- Every employee can view their own E-ID profile and complete a ONE-TIME self edit
  (add photo + personal details) right after first login; after that it locks and only
  admin/HR can change it.
"""
import io
import os
import zipfile

from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from .. import models, audit
from ..database import get_db
from ..deps import get_current_user, require_roles
from ..storage import save_photo, resolve as resolve_photo, read_bytes

router = APIRouter(prefix="/api/manpower", tags=["manpower"])

HR_ROLES = ("admin", "headoffice", "hr")

# Fields an employee may set during their one-time self edit.
SELF_EDITABLE = {
    "phone", "gender", "dob", "blood_group", "marital_status", "emergency_contact",
    "emergency_name", "emergency_relation", "current_address", "aadhar_number",
    "pan_number", "bank_holder", "bank_account", "ifsc_code", "bank_name",
    "aadhar_address", "photo_url", "address",
}

# --- Profile change-request flow -------------------------------------------------
# After the one-time lock, an employee can REQUEST a change to any of their profile
# fields with the actual new value; HR / Admin review it and, on approval, the value
# is applied to their profile.
CR_APPROVERS = ("admin", "hr")           # who reviews change requests (HR/Admin only)

# Any profile field the employee may request a change to, with a human label. Excludes
# system-managed identity (id, emp_code, role, is_active, profile_completed, team-lead hat).
REQUESTABLE_FIELDS = {
    "name": "Full name", "email": "Login email", "phone": "Phone",
    "designation": "Designation", "location": "Location", "branch": "Branch",
    "gender": "Gender", "dob": "Date of birth", "joining_date": "Date of joining",
    "blood_group": "Blood group", "marital_status": "Marital status",
    "emergency_contact": "Emergency contact no.", "emergency_name": "Emergency contact name",
    "emergency_relation": "Emergency contact relation", "aadhar_number": "Aadhaar number",
    "pan_number": "PAN number", "bank_holder": "Bank account holder",
    "bank_account": "Bank account number", "ifsc_code": "IFSC code", "bank_name": "Bank name",
    "current_address": "Current address", "aadhar_address": "Aadhaar address",
    "rent_own": "Rent / Own", "hr_ref": "HR reference", "ctc": "CTC",
}
_CR_DATE_FIELDS = ("dob", "joining_date")


def _cr_display(u: models.User, field: str):
    """Current value of a requestable field as a plain string for the diff view."""
    v = getattr(u, field, None)
    if field in _CR_DATE_FIELDS and v:
        try:
            return v.isoformat()
        except Exception:
            return str(v)
    return "" if v is None else str(v)


def _pcr(r: "models.ProfileChangeRequest", names: dict) -> dict:
    nm = names.get(r.user_id, (None, None, None))
    rv = names.get(r.reviewed_by, (None, None, None)) if r.reviewed_by else (None, None, None)
    return {
        "id": r.id, "user_id": r.user_id, "user_name": nm[0], "user_code": nm[2],
        "user_branch": nm[1], "field": r.field, "field_label": r.field_label,
        "old_value": r.old_value, "new_value": r.new_value, "note": r.note,
        "status": r.status, "review_note": r.review_note,
        "reviewed_by": r.reviewed_by, "reviewer_name": rv[0],
        "reviewed_at": r.reviewed_at.isoformat() if r.reviewed_at else None,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


def _cr_names(db: Session) -> dict:
    return {u.id: (u.name, u.branch, u.emp_code) for u in db.query(models.User).all()}


def _emp(u: models.User) -> dict:
    return {
        "id": u.id, "emp_code": u.emp_code, "hr_ref": u.hr_ref, "name": u.name,
        "email": u.email, "phone": u.phone, "role": u.role, "designation": u.designation,
        "location": u.location, "branch": u.branch, "gender": u.gender,
        "dob": u.dob.isoformat() if u.dob else None,
        "joining_date": u.joining_date.isoformat() if u.joining_date else None,
        "blood_group": u.blood_group, "marital_status": u.marital_status,
        "emergency_contact": u.emergency_contact, "emergency_name": u.emergency_name,
        "emergency_relation": u.emergency_relation, "dra_status": u.dra_status,
        "pvc_status": u.pvc_status, "aadhar_number": u.aadhar_number, "pan_number": u.pan_number,
        "bank_holder": u.bank_holder, "bank_account": u.bank_account, "ifsc_code": u.ifsc_code,
        "bank_name": u.bank_name, "current_address": u.current_address,
        "aadhar_address": u.aadhar_address, "rent_own": u.rent_own, "ctc": u.ctc,
        "photo_url": resolve_photo(u.photo_url), "is_active": u.is_active,
        "blocked_reason": u.blocked_reason,
        "blocked_at": u.blocked_at.isoformat() if u.blocked_at else None,
        "profile_completed": bool(u.profile_completed),
        "also_team_lead": bool(u.also_team_lead), "tl_emp_code": u.tl_emp_code,
        "all_roles": _all_roles(u), "all_ids": _all_ids(u),
    }


_ROLE_LABELS = {"telecaller": "Tele-calling Agent", "fos": "Field Agent", "teamlead": "Team Lead",
                "manager": "Collections Manager", "admin": "Administrator", "headoffice": "Head Office",
                "backend": "Back-office Official", "hr": "HR", "it": "IT", "staff": "Staff"}


def _all_roles(u) -> str:
    """Both roles for a dual-role user, e.g. 'Field Agent + Team Lead'."""
    label = _ROLE_LABELS.get(u.role, (u.role or "").title())
    if getattr(u, "also_team_lead", False) and u.role != "teamlead":
        return f"{label} + Team Lead"
    return label


def _all_ids(u) -> str:
    """Both IDs for a dual-role user, e.g. 'FO012 / TL014'."""
    ids = [u.emp_code] if u.emp_code else []
    if getattr(u, "also_team_lead", False) and getattr(u, "tl_emp_code", None):
        ids.append(u.tl_emp_code)
    return " / ".join(ids)


@router.get("/filters")
def filters(db: Session = Depends(get_db), user: models.User = Depends(require_roles(*HR_ROLES))):
    us = db.query(models.User.role, models.User.location).all()
    roles = sorted({r for r, _ in us if r})
    locations = sorted({(loc or "").strip() for _, loc in us if loc and loc.strip()})
    return {"roles": roles, "locations": locations}


@router.get("")
def list_manpower(role: str | None = None, location: str | None = None, q: str | None = None,
                  db: Session = Depends(get_db), user: models.User = Depends(require_roles(*HR_ROLES))):
    """The full employee directory with role / location / text filters."""
    query = db.query(models.User).filter(models.User.role != "techsupport")   # hidden support role
    if role == "teamlead":
        # Include dual-role staff (a caller/FOS who also wears the team-lead hat) in the
        # team-leads directory, so their team is reachable too.
        from sqlalchemy import or_ as _or
        query = query.filter(_or(models.User.role == "teamlead", models.User.also_team_lead.is_(True)))
    elif role:
        query = query.filter(models.User.role == role)
    if location:
        query = query.filter(func_lower(models.User.location) == location.strip().lower())
    rows = query.order_by(models.User.name).all()
    out = [_emp(u) for u in rows]
    if q:
        s = q.lower().strip()
        def _hit(e):
            return any(s in (str(e.get(k) or "")).lower() for k in
                       ("name", "emp_code", "email", "phone", "designation",
                        "location", "branch", "hr_ref", "role"))
        out = [e for e in out if _hit(e)]
    return out


@router.get("/download")
def download(role: str | None = None, location: str | None = None,
             db: Session = Depends(get_db), user: models.User = Depends(require_roles(*HR_ROLES))):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    query = db.query(models.User)
    if role:
        query = query.filter(models.User.role == role)
    if location:
        query = query.filter(func_lower(models.User.location) == location.strip().lower())
    rows = query.order_by(models.User.role, models.User.name).all()

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Manpower"
    cols = [("emp_code", "Emp Code"), ("tl_emp_code", "Team Lead ID"), ("all_roles", "Roles"),
            ("hr_ref", "HR Ref"), ("name", "Name"),
            ("role", "Primary Role"), ("designation", "Designation"), ("location", "Location"),
            ("branch", "Branch"), ("phone", "Phone"), ("email", "Email"), ("gender", "Gender"),
            ("dob", "DOB"), ("joining_date", "DOJ"), ("blood_group", "Blood"),
            ("marital_status", "Marital"), ("emergency_contact", "Emergency No"),
            ("emergency_name", "Emergency Name"), ("emergency_relation", "Relation"),
            ("aadhar_number", "Aadhaar"), ("pan_number", "PAN"), ("bank_holder", "Bank Holder"),
            ("bank_account", "Account"), ("ifsc_code", "IFSC"), ("bank_name", "Bank"),
            ("current_address", "Current Address"), ("rent_own", "Rent/Own")]
    ws.append([c[1] for c in cols])
    for i in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=i); cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for u in rows:
        e = _emp(u)
        ws.append([e.get(c[0]) for c in cols])
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=SSDE_manpower.xlsx"})


# Fields an HR user can set when adding an employee (mirrors the manpower sheet).
_ADD_FIELDS = (
    "phone", "designation", "location", "branch", "hr_ref", "gender", "blood_group",
    "marital_status", "emergency_contact", "emergency_name", "emergency_relation",
    "aadhar_number", "pan_number", "bank_holder", "bank_account", "ifsc_code", "bank_name",
    "current_address", "aadhar_address", "employment_type", "photo_url", "ctc",
)
_STARTER_PASSWORD = "Ssd@2026"


@router.post("")
def add_employee(body: dict = Body(...), db: Session = Depends(get_db),
                 user: models.User = Depends(require_roles("admin", "headoffice", "hr"))):
    """Add a new employee to the directory with full HR details. Creates a login with a
    starter password (they set their own on first sign-in) and a role-wise emp code."""
    from ..security import hash_password
    from .users import generate_emp_code
    from datetime import datetime as _dt

    name = (body.get("name") or "").strip()
    email = (body.get("email") or "").strip().lower()
    role = (body.get("role") or "").strip()
    if not name or not email or not role:
        raise HTTPException(status_code=400, detail="Name, login email and role are required")
    if role == "admin" and user.role != "admin":
        raise HTTPException(status_code=403,
                            detail="Only an administrator can create another Administrator.")
    if db.query(models.User).filter(models.User.email == email).first():
        raise HTTPException(status_code=400, detail="That login email is already registered")

    u = models.User(
        name=name, email=email, role=role,
        emp_code=generate_emp_code(db, role),
        hashed_password=hash_password(body.get("password") or _STARTER_PASSWORD),
        must_change_password=True, profile_completed=True, is_active=True,
    )
    for k in _ADD_FIELDS:
        v = body.get(k)
        if v not in (None, ""):
            setattr(u, k, v)
    for dk in ("dob", "joining_date"):
        dv = body.get(dk)
        if dv:
            try:
                setattr(u, dk, _dt.strptime(str(dv)[:10], "%Y-%m-%d").date())
            except ValueError:
                pass
    db.add(u)
    db.flush()
    audit.record(db, user, "staff_create", None, entity_type="staff", target_user_id=u.id,
                 detail=f"Added employee {name} ({role}) — {u.emp_code}")
    db.commit()
    db.refresh(u)
    return _emp(u)


@router.get("/me")
def my_profile(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """The signed-in employee's own profile / E-ID data."""
    return _emp(user)


@router.post("/me/photo")
async def upload_my_photo(file: UploadFile = File(...), db: Session = Depends(get_db),
                         user: models.User = Depends(get_current_user)):
    """Upload / replace the signed-in employee's profile photo. Allowed ANY time (not
    subject to the one-time profile lock), so staff can update their picture whenever."""
    ct = (file.content_type or "").lower()
    if not ct.startswith("image/"):
        raise HTTPException(status_code=400, detail="Please upload an image file")
    content = await file.read()
    if len(content) > 8 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Image too large (max 8 MB)")
    ref = save_photo(content, file.filename or "photo.jpg", ct, folder="profiles")
    user.photo_url = ref
    audit.record(db, user, "profile_update", None, entity_type="staff", detail="Updated profile photo")
    db.commit()
    return {"ok": True, "photo_url": resolve_photo(ref)}


@router.patch("/me")
def update_my_profile(body: dict = Body(...), db: Session = Depends(get_db),
                      user: models.User = Depends(get_current_user)):
    """One-time self edit right after first login. Locks once completed (admin/HR can
    still edit via user management). Admin/HR are not limited by the one-time lock."""
    is_hr = user.role in HR_ROLES
    if user.profile_completed and not is_hr:
        raise HTTPException(status_code=403,
                            detail="Your profile is locked. Submit a change request for HR/Admin approval.")
    changed = 0
    for k, v in body.items():
        if k not in SELF_EDITABLE:
            continue
        if k in ("dob",) and v:
            from datetime import datetime as _dt
            try:
                v = _dt.strptime(str(v)[:10], "%Y-%m-%d").date()
            except ValueError:
                continue
        setattr(user, k, v)
        changed += 1
    if not is_hr:
        user.profile_completed = True    # lock after the one-time completion
    audit.record(db, user, "profile_update", None, entity_type="staff",
                 detail=f"{'HR' if is_hr else 'Self'} profile update ({changed} fields)")
    db.commit()
    return {"ok": True, "profile_completed": bool(user.profile_completed)}


# --- Profile change requests (locked employees) --------------------------------
@router.get("/me/change-fields")
def my_change_fields(user: models.User = Depends(get_current_user)):
    """The fields an employee may request a change to, each with its current value."""
    return [{"field": f, "label": lbl, "current": _cr_display(user, f)}
            for f, lbl in REQUESTABLE_FIELDS.items()]


@router.post("/me/change-request")
def submit_change_request(body: dict = Body(...), db: Session = Depends(get_db),
                          user: models.User = Depends(get_current_user)):
    """Employee asks HR/Admin to change one profile field to a specific new value."""
    field = (body.get("field") or "").strip()
    if field not in REQUESTABLE_FIELDS:
        raise HTTPException(status_code=400, detail="That field cannot be changed by request.")
    new_value = body.get("value")
    new_value = "" if new_value is None else str(new_value).strip()
    if not new_value:
        raise HTTPException(status_code=400, detail="Enter the new value you want.")
    old_value = _cr_display(user, field)
    if new_value == old_value:
        raise HTTPException(status_code=400, detail="That value is the same as your current one.")
    # One open request per field — replace any earlier pending one for the same field.
    db.query(models.ProfileChangeRequest).filter(
        models.ProfileChangeRequest.user_id == user.id,
        models.ProfileChangeRequest.field == field,
        models.ProfileChangeRequest.status == "pending").update(
        {models.ProfileChangeRequest.status: "superseded"})
    r = models.ProfileChangeRequest(
        user_id=user.id, field=field, field_label=REQUESTABLE_FIELDS[field],
        old_value=old_value, new_value=new_value,
        note=(body.get("note") or "").strip()[:300] or None, status="pending")
    db.add(r)
    db.flush()
    # Alert HR / Admin.
    from .notifications import push
    for appr in db.query(models.User).filter(models.User.role.in_(CR_APPROVERS),
                                             models.User.is_active.is_(True)).all():
        push(db, appr.id, f"Profile change request — {user.name}",
             f"{REQUESTABLE_FIELDS[field]}: “{old_value or '—'}” → “{new_value}”",
             ntype="profile_request", by_name=user.name)
    audit.record(db, user, "profile_update", None, entity_type="staff",
                 detail=f"Requested change: {REQUESTABLE_FIELDS[field]} → {new_value}")
    db.commit()
    db.refresh(r)
    return _pcr(r, _cr_names(db))


@router.get("/me/change-requests")
def my_change_requests(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """The signed-in employee's own change requests, newest first."""
    rows = (db.query(models.ProfileChangeRequest)
            .filter(models.ProfileChangeRequest.user_id == user.id)
            .order_by(models.ProfileChangeRequest.created_at.desc()).all())
    names = _cr_names(db)
    return [_pcr(r, names) for r in rows]


@router.get("/change-requests")
def list_change_requests(status: str | None = None, q: str | None = None,
                         db: Session = Depends(get_db),
                         user: models.User = Depends(require_roles(*CR_APPROVERS))):
    """HR / Admin: all employees' profile change requests, filterable by status + text."""
    query = db.query(models.ProfileChangeRequest)
    if status and status != "all":
        query = query.filter(models.ProfileChangeRequest.status == status)
    rows = query.order_by(models.ProfileChangeRequest.created_at.desc()).all()
    names = _cr_names(db)
    out = [_pcr(r, names) for r in rows]
    if q:
        s = q.lower().strip()
        out = [e for e in out if any(s in (str(e.get(k) or "")).lower() for k in
               ("user_name", "user_code", "user_branch", "field_label", "new_value", "old_value"))]
    return out


@router.get("/change-requests/pending-count")
def change_requests_pending_count(db: Session = Depends(get_db),
                                  user: models.User = Depends(require_roles(*CR_APPROVERS))):
    n = db.query(models.ProfileChangeRequest).filter(
        models.ProfileChangeRequest.status == "pending").count()
    return {"pending": n}


@router.post("/change-requests/{req_id}/{decision}")
def decide_change_request(req_id: int, decision: str, body: dict = Body(default={}),
                          db: Session = Depends(get_db),
                          user: models.User = Depends(require_roles(*CR_APPROVERS))):
    """HR / Admin approve (applies the value) or reject a profile change request."""
    if decision not in ("approve", "reject"):
        raise HTTPException(status_code=400, detail="decision must be approve or reject")
    r = db.query(models.ProfileChangeRequest).filter(
        models.ProfileChangeRequest.id == req_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Request not found")
    if r.status != "pending":
        raise HTTPException(status_code=400, detail=f"This request is already {r.status}.")
    emp = db.query(models.User).filter(models.User.id == r.user_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")

    if decision == "approve":
        field, val = r.field, r.new_value
        if field == "email":
            new_email = (val or "").strip().lower()
            if new_email and db.query(models.User).filter(
                    models.User.email == new_email, models.User.id != emp.id).first():
                raise HTTPException(status_code=400, detail="That login email is already registered.")
            emp.email = new_email
        elif field in _CR_DATE_FIELDS:
            from datetime import datetime as _dt
            try:
                setattr(emp, field, _dt.strptime(str(val)[:10], "%Y-%m-%d").date() if val else None)
            except ValueError:
                raise HTTPException(status_code=400, detail="Approved value is not a valid date (YYYY-MM-DD).")
        else:
            setattr(emp, field, val if val not in ("",) else None)
        r.status = "approved"
    else:
        r.status = "rejected"
    from datetime import datetime as _dtn, timezone as _tz
    r.reviewed_by = user.id
    r.reviewed_at = _dtn.now(_tz.utc)
    r.review_note = (body.get("review_note") or "").strip()[:300] or None

    from .notifications import push
    verb = "approved" if decision == "approve" else "rejected"
    push(db, emp.id, f"Profile change {verb}",
         f"{r.field_label} → “{r.new_value}” was {verb}"
         + (f" · {r.review_note}" if r.review_note else ""),
         ntype="profile_request", by_name=user.name)
    audit.record(db, user, "staff_update", None, entity_type="staff", target_user_id=emp.id,
                 detail=f"{verb.title()} profile change: {r.field_label} → {r.new_value}")
    db.commit()
    db.refresh(r)
    return _pcr(r, _cr_names(db))


# Registered AFTER the /me routes so "/me" is never captured as an emp_id.
@router.patch("/{emp_id}")
def edit_employee(emp_id: int, body: dict = Body(...), db: Session = Depends(get_db),
                  user: models.User = Depends(require_roles("admin", "headoffice", "hr"))):
    """HR / admin / head office edits any employee's directory details. Only the fields sent
    are changed. Can also reset the person's password (forces a change on next login)."""
    from datetime import datetime as _dt
    u = db.query(models.User).filter(models.User.id == emp_id).first()
    if not u:
        raise HTTPException(status_code=404, detail="Employee not found")

    new_email = (body.get("email") or "").strip().lower()
    if new_email and new_email != (u.email or "").lower():
        if db.query(models.User).filter(models.User.email == new_email,
                                        models.User.id != u.id).first():
            raise HTTPException(status_code=400, detail="That login email is already registered")
        u.email = new_email
    if (body.get("name") or "").strip():
        u.name = body["name"].strip()
    old_role = u.role
    new_role = (body.get("role") or "").strip()
    if new_role:
        # Only an admin may grant the admin role, or change an existing admin's role.
        if (new_role == "admin" or old_role == "admin") and user.role != "admin":
            raise HTTPException(status_code=403,
                                detail="Only an administrator can assign or change the Administrator role.")
        u.role = new_role
    # Optional: when the role changes, regenerate the emp code so its prefix matches the new
    # role (e.g. FO001 → TL003). Off by default — existing sheet references keep the old code.
    if body.get("regen_code") and u.role != old_role:
        from .users import generate_emp_code
        u.emp_code = generate_emp_code(db, u.role)
    # Dual role: grant/revoke the team-lead hat on a caller/FOS. Granting keeps their primary
    # role + emp_code and issues a second team-lead ID (tl_emp_code) if they don't have one.
    if "also_team_lead" in body:
        grant = bool(body["also_team_lead"])
        u.also_team_lead = grant
        if grant and u.role != "teamlead" and not u.tl_emp_code:
            from .users import generate_emp_code
            u.tl_emp_code = generate_emp_code(db, "teamlead")
    if "is_active" in body:
        active = bool(body["is_active"])
        if active and not u.is_active:            # unblocking → clear the block record
            u.blocked_reason = None
            u.blocked_at = None
            u.blocked_by = None
        elif not active:                          # blocking / removing → stamp who + why + when
            u.blocked_reason = (body.get("block_reason") or u.blocked_reason or "").strip()[:200] or None
            u.blocked_at = _dt.utcnow()
            u.blocked_by = user.id
        u.is_active = active
        # A blocked employee's cases stay assigned (history preserved) but they can no longer log in
        # and are filtered out of the active FOS/caller pickers everywhere.
    for k in _ADD_FIELDS:
        if k in body:
            v = body.get(k)
            setattr(u, k, v if v not in ("",) else None)
    for dk in ("dob", "joining_date"):
        if dk in body:
            dv = body.get(dk)
            if dv:
                try:
                    setattr(u, dk, _dt.strptime(str(dv)[:10], "%Y-%m-%d").date())
                except ValueError:
                    pass
            else:
                setattr(u, dk, None)
    if body.get("password"):
        from ..security import hash_password
        u.hashed_password = hash_password(body["password"])
        u.must_change_password = True
    audit.record(db, user, "staff_update", None, entity_type="staff", target_user_id=u.id,
                 detail=f"Edited employee {u.name} ({u.emp_code})")
    db.commit()
    db.refresh(u)
    return _emp(u)


# tiny local helper to avoid importing func at module top for one use
def func_lower(col):
    from sqlalchemy import func
    return func.lower(col)


# ---------------------------------------------------------------------------
# HR document vault — upload / list / download / zip per employee.
# ---------------------------------------------------------------------------
DOC_TYPES = [
    ("pan", "PAN card"),
    ("aadhaar", "Aadhaar card"),
    ("photo", "Photo"),
    ("signature", "Signature (white paper)"),
    ("pvc", "PVC"),
    ("dra", "DRA certificate"),
    ("cibil", "CIBIL report (Paisabazaar)"),
    ("bank_details", "Bank account details"),
    ("reference", "Reference contact details"),
    ("whatsapp", "WhatsApp no. (not PhonePe-linked)"),
    ("email", "Email ID proof"),
]
# "other" is a catch-all bucket for folder uploads whose filename didn't map to a known type.
# Unlike the 11 fixed types it may hold MANY files (they don't replace each other).
_DOC_KEYS = {k for k, _ in DOC_TYPES} | {"other"}


@router.get("/doc-types")
def doc_types(user: models.User = Depends(require_roles(*HR_ROLES))):
    return [{"key": k, "label": lbl} for k, lbl in DOC_TYPES]


@router.get("/{emp_id}/documents")
def list_documents(emp_id: int, db: Session = Depends(get_db),
                   user: models.User = Depends(require_roles(*HR_ROLES))):
    rows = (db.query(models.EmployeeDocument)
            .filter(models.EmployeeDocument.user_id == emp_id)
            .order_by(models.EmployeeDocument.uploaded_at.desc()).all())
    latest = {}
    others = []
    for d in rows:                       # keep the most-recent per FIXED type; keep ALL "other"
        if d.doc_type == "other":
            others.append(d)
        else:
            latest.setdefault(d.doc_type, d)

    def _ser(d):
        return {"id": d.id, "doc_type": d.doc_type, "filename": d.filename,
                "url": resolve_photo(d.ref),
                "uploaded_at": d.uploaded_at.isoformat() if d.uploaded_at else None}
    return [_ser(d) for d in latest.values()] + [_ser(d) for d in others]


@router.post("/{emp_id}/documents")
async def upload_document(emp_id: int, doc_type: str = Form(...), file: UploadFile = File(...),
                          db: Session = Depends(get_db),
                          user: models.User = Depends(require_roles(*HR_ROLES))):
    emp = db.query(models.User).filter(models.User.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    if doc_type not in _DOC_KEYS:
        raise HTTPException(status_code=400, detail="Unknown document type")
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Empty file")
    ref = save_photo(content, filename=file.filename or doc_type,
                     content_type=file.content_type or "application/octet-stream", folder="documents")
    # Replace any prior file of this type (keep one current per FIXED type). "other" accumulates.
    if doc_type != "other":
        db.query(models.EmployeeDocument).filter(
            models.EmployeeDocument.user_id == emp_id,
            models.EmployeeDocument.doc_type == doc_type).delete()
    d = models.EmployeeDocument(user_id=emp_id, doc_type=doc_type, filename=file.filename,
                                ref=ref, content_type=file.content_type, uploaded_by=user.id)
    db.add(d)
    audit.record(db, user, "staff_update", None, entity_type="staff", target_user_id=emp_id,
                 detail=f"Uploaded document '{doc_type}' for {emp.name}")
    db.commit()
    db.refresh(d)
    return {"id": d.id, "doc_type": d.doc_type, "filename": d.filename, "url": resolve_photo(d.ref)}


@router.post("/{emp_id}/documents/bulk")
async def upload_documents_bulk(emp_id: int,
                                files: list[UploadFile] = File(...),
                                doc_types: list[str] = Form(...),
                                db: Session = Depends(get_db),
                                user: models.User = Depends(require_roles(*HR_ROLES))):
    """Upload a whole folder for one employee. `files` and `doc_types` are parallel lists —
    the caller maps each file to a document type (auto-detected from the filename, HR-confirmed).
    Unknown types fall into 'other'; a type of 'skip' drops that file."""
    emp = db.query(models.User).filter(models.User.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    if len(files) != len(doc_types):
        raise HTTPException(status_code=400, detail="files and doc_types count mismatch")
    saved = []
    cleared: set[str] = set()
    for f, raw_dt in zip(files, doc_types):
        dt = (raw_dt or "").strip()
        if dt == "skip" or not dt:
            continue
        if dt not in _DOC_KEYS:
            dt = "other"
        content = await f.read()
        if not content:
            continue
        ref = save_photo(content, filename=f.filename or dt,
                         content_type=f.content_type or "application/octet-stream", folder="documents")
        if dt != "other" and dt not in cleared:      # replace prior once per fixed type
            db.query(models.EmployeeDocument).filter(
                models.EmployeeDocument.user_id == emp_id,
                models.EmployeeDocument.doc_type == dt).delete()
            cleared.add(dt)
        db.add(models.EmployeeDocument(user_id=emp_id, doc_type=dt, filename=f.filename,
                                       ref=ref, content_type=f.content_type, uploaded_by=user.id))
        saved.append({"doc_type": dt, "filename": f.filename})
    audit.record(db, user, "staff_update", None, entity_type="staff", target_user_id=emp_id,
                 detail=f"Bulk-uploaded {len(saved)} document(s) for {emp.name}")
    db.commit()
    return {"uploaded": len(saved), "documents": saved}


@router.get("/{emp_id}/documents.zip")
def documents_zip(emp_id: int, db: Session = Depends(get_db),
                  user: models.User = Depends(require_roles(*HR_ROLES))):
    emp = db.query(models.User).filter(models.User.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    rows = db.query(models.EmployeeDocument).filter(models.EmployeeDocument.user_id == emp_id).all()
    buf = io.BytesIO()
    seen: dict[str, int] = {}
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for d in rows:
            data = read_bytes(d.ref)
            if data is None:
                continue
            ext = os.path.splitext(d.filename or "")[1] or ""
            # Fixed types are named by their type; "other" keeps its original filename.
            base = d.doc_type if d.doc_type != "other" else (os.path.splitext(d.filename or "other")[0] or "other")
            name = f"{base}{ext}"
            if name in seen:                         # avoid collisions when a type has >1 file
                seen[name] += 1
                name = f"{base}_{seen[name]}{ext}"
            else:
                seen[name] = 0
            z.writestr(name, data)
    buf.seek(0)
    safe = "".join(ch for ch in (emp.name or f"emp{emp_id}")
                   if ch.isalnum() or ch in " _-").strip().replace(" ", "_") or f"emp{emp_id}"
    return StreamingResponse(buf, media_type="application/zip",
                             headers={"Content-Disposition": f'attachment; filename="{safe}.zip"'})


@router.get("/{emp_id}/documents/{doc_id}/download")
def download_document(emp_id: int, doc_id: int, db: Session = Depends(get_db),
                      user: models.User = Depends(require_roles(*HR_ROLES))):
    d = (db.query(models.EmployeeDocument)
         .filter(models.EmployeeDocument.id == doc_id, models.EmployeeDocument.user_id == emp_id).first())
    if not d:
        raise HTTPException(status_code=404, detail="Document not found")
    data = read_bytes(d.ref)
    if data is None:
        raise HTTPException(status_code=404, detail="File missing from storage")
    return StreamingResponse(io.BytesIO(data), media_type=d.content_type or "application/octet-stream",
                             headers={"Content-Disposition": f'attachment; filename="{d.filename or d.doc_type}"'})


@router.delete("/{emp_id}/documents/{doc_id}")
def delete_document(emp_id: int, doc_id: int, db: Session = Depends(get_db),
                    user: models.User = Depends(require_roles(*HR_ROLES))):
    n = (db.query(models.EmployeeDocument)
         .filter(models.EmployeeDocument.id == doc_id, models.EmployeeDocument.user_id == emp_id)
         .delete())
    db.commit()
    return {"deleted": n}


# ---------------------------------------------------------------------------
# Offer-letter generator — preview (editable HTML) + email to the candidate.
# ---------------------------------------------------------------------------
def _esc(x) -> str:
    return str("" if x is None else x).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _fmt_amount(v):
    """Nicely format a rupee figure; leave placeholders / text as-is."""
    try:
        n = float(str(v).replace(",", "").replace("₹", "").replace("/-", "").strip())
        return "₹ {:,.0f}/-".format(n)
    except (TypeError, ValueError):
        return str(v)


def _company(d: dict) -> dict:
    """SSD company block for the letters (each field can be overridden from the request)."""
    return {
        "name": d.get("company") or "SRI SAI DHANADA ENTERPRISES",
        "address": d.get("company_address")
        or "39-11-5, 4th Floor, Axis Bank Upstairs, Opp Union Bank, Murali Nagar, Visakhapatnam - 530007",
        "hr_name": d.get("hr_name") or "Nambala Santhi Kumari",
        "hr_email": d.get("hr_email") or "ssdenterpriseshr@gmail.com",
        "signatory": d.get("signatory") or "S Govind Rao",
        "signatory_title": d.get("signatory_title") or "Managing Partner",
    }


_NAVY = "#0B234F"
_GOLD = "#C7A24A"
_LOGO_URI = None


def _logo_data_uri() -> str:
    """The SSD gold-coin logo (from the company brochure) as a base64 data URI so it embeds
    directly in the letter HTML — works in print, download, and email without a hosted URL."""
    global _LOGO_URI
    if _LOGO_URI is None:
        import base64
        p = os.path.join(os.path.dirname(__file__), "..", "assets", "ssd_logo.png")
        try:
            with open(p, "rb") as f:
                _LOGO_URI = "data:image/png;base64," + base64.b64encode(f.read()).decode()
        except Exception:
            _LOGO_URI = ""
    return _LOGO_URI


def _letterhead(c: dict, subtitle: str) -> str:
    logo = _logo_data_uri()
    logo_html = (f'<img src="{logo}" alt="SSD" style="height:78px;width:auto;display:block;margin:0 auto 8px" />'
                 if logo else "")
    return f"""<div style="background:{_NAVY};border:1px solid {_GOLD};border-radius:12px;padding:22px 20px 18px;text-align:center;margin-bottom:22px">
    {logo_html}
    <div style="font-family:Georgia,'Times New Roman',serif;font-size:23px;font-weight:700;letter-spacing:1px;color:#F4E9C9">{_esc(c['name'])}</div>
    <div style="font-size:10px;letter-spacing:3px;color:{_GOLD};text-transform:uppercase;margin-top:5px">Debt Recovery Agency &amp; Management</div>
    <div style="font-size:11px;color:#c3cde0;margin-top:7px">{_esc(c['address'])}</div>
    <div style="display:inline-block;margin-top:11px;padding:3px 16px;border:1px solid {_GOLD};color:{_GOLD};border-radius:999px;font-size:10.5px;font-weight:700;letter-spacing:2px;text-transform:uppercase">{_esc(subtitle)}</div>
  </div>"""


def _sign_block(c: dict, name: str) -> str:
    return f"""<table style="width:100%;margin-top:40px;font-size:13px"><tr>
    <td style="width:50%;vertical-align:top">
      <div style="height:34px"></div>
      <div style="border-top:2px solid {_GOLD};width:230px;padding-top:4px;color:{_NAVY}">For <b>{_esc(c['name'])}</b></div>
      <div>{_esc(c['signatory'])}</div>
      <div style="color:#6b7280">{_esc(c['signatory_title'])}</div>
    </td>
    <td style="width:50%;vertical-align:top">
      <div style="height:34px"></div>
      <div style="border-top:2px solid {_GOLD};width:230px;padding-top:4px;color:{_NAVY}">Accepted by <b>{_esc(name)}</b></div>
      <div style="color:#6b7280">Signature &amp; Date</div>
    </td>
  </tr></table>"""


_H = "color:#0B234F;font-size:15px;margin:20px 0 8px;font-weight:700;border-left:4px solid #C7A24A;padding-left:8px"
_WRAP = ("font-family:'Segoe UI',Calibri,Arial,sans-serif;color:#1f2937;max-width:760px;"
         "margin:0 auto;line-height:1.65;font-size:14px;background:#fff;padding:26px 30px;"
         "border:3px double #C7A24A;border-radius:12px;box-shadow:0 6px 24px rgba(11,35,79,.10)")


def _offer_letter_html(emp, d: dict) -> str:
    from datetime import date as _date
    c = _company(d)
    today = d.get("date") or _date.today().strftime("%d/%m/%Y")
    name = d.get("name") or emp.name or "[Employee Name]"
    designation = d.get("designation") or emp.designation or "[Designation]"
    reporting = d.get("reporting_to") or c["hr_name"]
    start = d.get("joining_date") or "[Start Date]"
    emp_type = d.get("employment_type") or "Full Time"
    schedule = d.get("work_schedule") or "9:30 AM to 7:30 PM"
    location = d.get("location") or emp.location or emp.branch or "In-Office"
    ctc = _fmt_amount(d.get("ctc") or emp.ctc or "[Annual CTC]")
    probation = d.get("probation_days") or "60"

    def row(k, v):
        return (f"<tr><td style='padding:5px 10px;color:#6b7280;width:180px'>{_esc(k)}</td>"
                f"<td style='padding:5px 10px;font-weight:600'>{_esc(v)}</td></tr>")
    details = row("Job Title", designation) + row("Reporting To", reporting) + row("Start Date", start) \
        + row("Employment Type", emp_type) + row("Work Schedule", schedule) + row("Job Location", location)
    terms = d.get("terms") or [
        "The employment is at-will — either the company or the employee may terminate the relationship at any time, with or without cause and with or without notice.",
        f"This offer does not constitute a contract or guarantee of continued employment until you have signed the employment agreement and any other required documents with “{c['name']}”.",
        f"During the probationary period of {probation} days, your performance will be evaluated to determine your suitability for the role.",
        "You will be required to sign a Confidentiality / Non-Compete Agreement after accepting this offer to protect the company's interests.",
    ]
    terms_html = "".join(f"<li style='margin-bottom:6px'>{_esc(t)}</li>" for t in terms)
    return f"""<div style="{_WRAP}">
  {_letterhead(c, 'Offer of Employment')}
  <div style="text-align:right;font-size:13px;color:#374151">Date: {_esc(today)}</div>
  <p>Dear <b>{_esc(name)}</b>,</p>
  <p>We are pleased to extend to you an offer of employment for the position of <b>{_esc(designation)}</b>
     at <b>{_esc(c['name'])}</b>. We are confident that your skills, experience, and dedication will make a
     valuable contribution to our organization. Kindly review the terms and conditions outlined in this offer
     letter carefully. Should you find them acceptable, please indicate your acceptance by signing and returning
     a copy of this letter within the stipulated time.</p>
  <p>We look forward to welcoming you to our team and anticipate a mutually rewarding professional association.</p>
  <h3 style="{_H}">Position Details</h3>
  <table style="border-collapse:collapse;width:100%;background:#FBF6E8;border:1px solid #E3D3A6;border-radius:8px;font-size:13.5px">{details}</table>
  <h3 style="{_H}">Compensation &amp; Benefits</h3>
  <p style="margin:6px 0"><b>Annual Salary Package (CTC): {_esc(ctc)}</b></p>
  <p style="font-size:13px;color:#374151">The above-mentioned salary is the total cost to the company and includes all
     payments made and benefits provided by the company, directly or indirectly, to or on your behalf, whether as salary or otherwise.</p>
  <h3 style="{_H}">Terms &amp; Conditions</h3>
  <ol style="font-size:13.5px;padding-left:20px">{terms_html}</ol>
  <h3 style="{_H}">Acceptance</h3>
  <p style="font-size:13.5px">This Letter of Offer contains the proposed terms and conditions of your employment with the
     Employer and is subject to the preparation and execution of a formal Contract of Employment. If you have any questions
     or require further information, please contact <b>{_esc(c['hr_name'])}</b> at
     <a href="mailto:{_esc(c['hr_email'])}">{_esc(c['hr_email'])}</a>.</p>
  <p style="font-size:13.5px">I, <b>{_esc(name)}</b>, accept and agree to the proposed terms of employment and request that the
     Employer prepares a formal contract of employment for execution.</p>
  {_sign_block(c, name)}
</div>"""


def _agreement_letter_html(emp, d: dict) -> str:
    from datetime import date as _date
    c = _company(d)
    today = d.get("date") or _date.today().strftime("%d/%m/%Y")
    name = d.get("name") or emp.name or "[Employee Name]"
    designation = d.get("designation") or emp.designation or "[Department / Designation]"
    reporting = d.get("reporting_to") or c["hr_name"]
    location = d.get("location") or emp.location or emp.branch or "Visakhapatnam"
    commence = d.get("joining_date") or "[Commencement Date]"
    address = d.get("address") or getattr(emp, "current_address", None) or getattr(emp, "address", None) or "[Employee Address]"
    ctc_raw = d.get("ctc") or emp.ctc or "360000"
    ctc = _fmt_amount(ctc_raw)
    probation_months = d.get("probation_months") or "3"
    notice = d.get("notice_days") or "30"
    salary = d.get("salary") or [
        {"component": "Basic Salary", "annual": ""},
        {"component": "House Rent Allowance (HRA)", "annual": ""},
        {"component": "Special Allowance", "annual": ""},
        {"component": "Bonus", "annual": ""},
        {"component": "Leave Travel Allowance", "annual": ""},
        {"component": "Commission", "annual": ""},
    ]
    sal_rows = "".join(
        f"<tr><td style='border:1px solid #cbd5e1;padding:6px 10px'>{_esc(r.get('component'))}</td>"
        f"<td style='border:1px solid #cbd5e1;padding:6px 10px;text-align:right'>{_esc(_fmt_amount(r.get('annual')) if r.get('annual') else '[____]')}</td></tr>"
        for r in salary)
    docs = [
        "Educational certificates with mark sheets (10th, 12th, Graduation, Post-Graduation, certifications).",
        "Latest salary slip and salary certificate from your last employer.",
        "Official relieving letter from your last employer.",
        "Service / experience certificates from all previous employers.",
        "An updated copy of your CV / resume.",
        "Form 16 or a taxable income statement from your previous employer.",
        "Four passport-size colour photographs.",
        "Valid passport and work permit (foreign nationals only).",
        "Proof of age (Birth Certificate or SSC memo).",
        "Proof of address (Aadhaar, Voter ID, or utility bill).",
        "A copy of your PAN card.",
    ]
    docs_html = "".join(f"<li style='margin-bottom:4px'>{_esc(x)}</li>" for x in docs)
    terms_b = [
        ("Term of Employment", "Your employment is intended to be for an indefinite period, subject to the termination clauses in this agreement and applicable Indian laws."),
        ("Outside Activities & Conflicts", "You must devote yourself exclusively to the company's business and may not take up other paid work or business without the company's written permission. You must keep company information confidential. The company may transfer you to any department, branch, or associated company as needed."),
        ("Termination", f"Resignation requires {notice} days' notice or salary in lieu of notice; the company may relieve you earlier. The company may terminate without cause on {notice} days' notice or salary in lieu, and immediately for cause (misconduct, fraud, theft, breach, unauthorised absence over 3 days, or insolvency). On exit you must return all company property and records."),
        ("Holidays and Leave", "General holidays are declared at the start of the calendar year. You are entitled to vacation and sick leave per company policy. Casual leave without notice is treated as leave against loss of pay; medical leave must be supported by a medical report."),
        ("Disclosure of Information", "You must disclose any information that may conflict with your employment or the company's interests, and notify any change in personal details in writing within three (3) days."),
        ("Adherence to Company Policy", "You agree to follow all company policies, directions, and orders issued from time to time."),
        ("Travel", f"Your primary work location is {location}. You may be required to travel within India or overseas as necessary for your duties."),
        ("Non-Solicitation", "During employment and for one year after, you agree not to solicit employees to leave or solicit customers for competing entities."),
        ("Assignment", "This agreement is personal to you and cannot be transferred by you; the company may assign it to its parents, subsidiaries, or affiliates."),
        ("Arbitration", "This agreement is governed by the laws of India and disputes will be settled under the Indian Arbitration and Conciliation Act, 1996."),
    ]
    terms_b_html = "".join(
        f"<li style='margin-bottom:8px'><b>{_esc(t)}</b><div style='color:#374151'>{_esc(v)}</div></li>"
        for t, v in terms_b)
    return f"""<div style="{_WRAP}">
  {_letterhead(c, 'Employment Agreement / Appointment Order')}
  <div style="text-align:right;font-size:13px;color:#374151">Date: {_esc(today)}</div>
  <p style="margin:2px 0">To,<br/><b>{_esc(name)}</b><br/><span style="color:#374151;font-size:13px;white-space:pre-line">{_esc(address)}</span></p>
  <p>Dear <b>{_esc(name)}</b>,</p>
  <p>We are pleased to offer you a position of <b>{_esc(designation)}</b> with <b>{_esc(c['name'])}</b>
     (hereinafter referred to as the &ldquo;Agency&rdquo;). This is a regular, full-time position based in
     <b>{_esc(location)}</b>. You will be reporting to <b>{_esc(reporting)}</b>. The terms of employment in this
     document and its annexures are confidential and must not be disclosed to third parties without the company's prior approval.</p>
  <h3 style="{_H}">1. Compensation</h3>
  <p style="font-size:13.5px">Your CTC will be <b>{_esc(ctc)}</b>. The break-up of your annual gross salary is in
     <b>Annexure&nbsp;A</b>. You will be entitled to benefits under applicable Indian labour and employment laws and eligible
     to participate in the company's employee benefit plans. The company may modify, amend, or terminate any benefit at any time.</p>
  <h3 style="{_H}">2. Terms &amp; Conditions of Employment</h3>
  <p style="font-size:13.5px">Your employment is governed by the terms in <b>Annexure&nbsp;B</b>.</p>
  <h3 style="{_H}">3. Commencement of Employment</h3>
  <p style="font-size:13.5px">You are required to commence employment on <b>{_esc(commence)}</b>. This offer is not valid beyond
     this date unless extended by the company in writing.</p>
  <h3 style="{_H}">4. Probation Period</h3>
  <p style="font-size:13.5px">You will undergo a probation evaluation for <b>{_esc(probation_months)} months</b> from the date of joining.
     On successful completion, your status is confirmed to a full-time employee. The salary during probation is per Annexure&nbsp;A.</p>
  <h3 style="{_H}">5. Document Submission Requirements</h3>
  <p style="font-size:13.5px">On your date of commencement, please report to complete joining formalities and submit the documents in <b>Annexure&nbsp;C</b>.</p>
  <h3 style="{_H}">6. Employment Invention Assignment Agreement</h3>
  <p style="font-size:13.5px">You will execute and be bound by an Employment Invention Assignment Agreement (<b>Annexure&nbsp;D</b>),
     which shall coexist with this Employment Agreement.</p>
  <h3 style="{_H}">7. Entire Agreement</h3>
  <p style="font-size:13.5px">This letter agreement (with its annexures) supersedes any prior agreements or representations and may only be
     modified by a written agreement signed by both parties. Once you accept this offer and join the Company, this letter will serve as
     your formal Appointment Order.</p>
  {_sign_block(c, name)}
  <div style="page-break-before:always;border-top:2px dashed #cbd5e1;margin-top:28px;padding-top:16px"></div>
  <h3 style="{_H}">Annexure A — Salary Structure</h3>
  <table style="border-collapse:collapse;width:100%;font-size:13.5px">
    <thead><tr style="background:#0B234F;color:#F4E9C9"><th style="border:1px solid #cbd5e1;padding:6px 10px;text-align:left">Particulars</th><th style="border:1px solid #cbd5e1;padding:6px 10px;text-align:right">INR / Annum</th></tr></thead>
    <tbody>{sal_rows}
      <tr style="background:#FBF6E8;font-weight:700;color:#0B234F"><td style="border:1px solid #cbd5e1;padding:6px 10px">Total Annual Salary (CTC)</td><td style="border:1px solid #cbd5e1;padding:6px 10px;text-align:right">{_esc(ctc)}</td></tr>
    </tbody>
  </table>
  <h3 style="{_H}">Annexure B — Terms &amp; Conditions</h3>
  <ol style="font-size:13.5px;padding-left:20px">{terms_b_html}</ol>
  <h3 style="{_H}">Annexure C — Documents to Submit on Joining</h3>
  <ol style="font-size:13.5px;padding-left:20px">{docs_html}</ol>
  <p style="font-size:12.5px;color:#6b7280">Note: Please carry all original documents for validation.</p>
  <h3 style="{_H}">Annexure D — Employment Invention Assignment</h3>
  <p style="font-size:13.5px">All inventions, works, and intellectual property created during the course of your employment shall
     belong to the Company. You agree to assign such rights to the Company and to sign any documents required to perfect that assignment.</p>
  <p style="font-size:13.5px;margin-top:14px"><b>Acceptance of Offer:</b> I have read and accept this offer of employment.</p>
  {_sign_block(c, name)}
</div>"""


def _send_email(to: str, subject: str, html: str, attachments=None):
    from ..config import get_settings
    s = get_settings()
    if not s.smtp_host or not s.smtp_user:
        raise HTTPException(status_code=400,
                            detail="Email is not configured. Set SMTP_HOST, SMTP_USER and SMTP_PASSWORD.")
    import smtplib
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from email.mime.application import MIMEApplication
    msg = MIMEMultipart("mixed")
    msg["Subject"] = subject
    msg["From"] = s.smtp_from or s.smtp_user
    msg["To"] = to
    alt = MIMEMultipart("alternative")
    alt.attach(MIMEText("Please open this email in an HTML-capable client.", "plain"))
    alt.attach(MIMEText(html, "html"))
    msg.attach(alt)
    for fn, data, _ctype in (attachments or []):
        part = MIMEApplication(data)
        part.add_header("Content-Disposition", "attachment", filename=fn)
        msg.attach(part)
    with smtplib.SMTP(s.smtp_host, s.smtp_port, timeout=20) as server:
        server.starttls()
        server.login(s.smtp_user, s.smtp_password)
        server.sendmail(msg["From"], [to], msg.as_string())


@router.post("/offer-letter")
def offer_letter_preview(body: dict = Body(...), db: Session = Depends(get_db),
                         user: models.User = Depends(require_roles(*HR_ROLES))):
    emp = db.query(models.User).filter(models.User.id == body.get("emp_id")).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    return {"html": _offer_letter_html(emp, body), "email": emp.email}


@router.post("/offer-letter/email")
def offer_letter_email(body: dict = Body(...), db: Session = Depends(get_db),
                       user: models.User = Depends(require_roles(*HR_ROLES))):
    emp = db.query(models.User).filter(models.User.id == body.get("emp_id")).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    to = (body.get("to") or emp.email or "").strip()
    if not to:
        raise HTTPException(status_code=400, detail="No recipient email — add the employee's email first")
    html = body.get("html") or _offer_letter_html(emp, body)
    subject = body.get("subject") or f"Offer of Employment — {emp.name}"
    safe = "".join(ch for ch in (emp.name or "candidate") if ch.isalnum() or ch in " _-").strip().replace(" ", "_")
    full = f"<html><body>{html}</body></html>"
    _send_email(to, subject, full, attachments=[(f"OfferLetter_{safe}.html", full.encode("utf-8"), "text/html")])
    audit.record(db, user, "staff_update", None, entity_type="staff", target_user_id=emp.id,
                 detail=f"Emailed offer letter to {to}")
    db.commit()
    return {"sent": True, "to": to}


@router.post("/agreement-letter")
def agreement_letter_preview(body: dict = Body(...), db: Session = Depends(get_db),
                             user: models.User = Depends(require_roles(*HR_ROLES))):
    emp = db.query(models.User).filter(models.User.id == body.get("emp_id")).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    return {"html": _agreement_letter_html(emp, body), "email": emp.email}


@router.post("/agreement-letter/email")
def agreement_letter_email(body: dict = Body(...), db: Session = Depends(get_db),
                           user: models.User = Depends(require_roles(*HR_ROLES))):
    emp = db.query(models.User).filter(models.User.id == body.get("emp_id")).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    to = (body.get("to") or emp.email or "").strip()
    if not to:
        raise HTTPException(status_code=400, detail="No recipient email — add the employee's email first")
    html = body.get("html") or _agreement_letter_html(emp, body)
    subject = body.get("subject") or f"Employment Agreement — {emp.name}"
    safe = "".join(ch for ch in (emp.name or "employee") if ch.isalnum() or ch in " _-").strip().replace(" ", "_")
    full = f"<html><body>{html}</body></html>"
    _send_email(to, subject, full, attachments=[(f"AgreementLetter_{safe}.html", full.encode("utf-8"), "text/html")])
    audit.record(db, user, "staff_update", None, entity_type="staff", target_user_id=emp.id,
                 detail=f"Emailed agreement letter to {to}")
    db.commit()
    return {"sent": True, "to": to}
