"""Reset the database and seed a realistic TEST dataset from the ICICI FR X-BKT sheet,
with FOS mapped to their real branch from the FOS-wise portfolio list.

WHAT IT DOES (idempotent — safe to re-run):
  1. WIPES every table.
  2. Creates the 5 branches (Hyderabad, Telangana, Visakhapatnam, Vijayawada, Kadapa),
     each with a manager. **Visakhapatnam is the head branch / HQ.**
  3. Creates a head-office user (cross-branch) at Visakhapatnam.
  4. Creates a FOS account per field agent in the FR sheet, mapped to the correct branch
     from `fos_portfolio.xlsx` (matched by phone). Emails are built from their name.
  5. Creates a telecaller per caller in the FR sheet — all under **Visakhapatnam** —
     with name-based emails.
  6. Imports the MAIN sheet as ICICI / FR / Credit Card cases, assigning each to its
     caller and FOS; a case's branch is the FOS's branch (the field owner).

Every login password is the same for testing: Test@1234
Emails are firstname.lastname@ssd.local (deduped).

RUN (from app/backend, venv active):  python seed_fr_test.py
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook                          # noqa: E402
from app.database import Base, SessionLocal, engine          # noqa: E402
from app import models                                        # noqa: E402
from app.excel_io import import_workbook, record_to_case_kwargs  # noqa: E402
from app.security import hash_password                        # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FR_FILE = os.path.join(HERE, "seed_data", "fr_xbkt_july26.xlsx")
PORTFOLIO_FILE = os.path.join(HERE, "seed_data", "fos_portfolio.xlsx")
BANK, PRODUCT, SEGMENT = "ICICI", "FR", "Credit Card"
BRANCHES = ["Hyderabad", "Telangana", "Visakhapatnam", "Vijayawada", "Kadapa"]
HEAD_BRANCH = "Visakhapatnam"
# Fallback branch for FOS present in the FR sheet but missing from the portfolio list,
# inferred from the area code in their sheet label.
AREA_BRANCH = {"VJW": "Vijayawada", "KDP": "Kadapa", "NALGONDA": "Telangana",
               "MEDAK": "Telangana", "PEDDAPALLI": "Telangana", "TIRUPATI": "Kadapa",
               "TS": "Telangana", "GTR": "Vijayawada"}
PW = hash_password("Test@1234")


def digits(s):
    d = re.sub(r"[^0-9]", "", str(s or ""))
    return d[-10:] if len(d) >= 10 else d


def slugify(name):
    s = re.sub(r"[^a-z0-9]+", ".", str(name or "").lower()).strip(".")
    return s or "user"


def load_portfolio():
    """phone(10) -> (proper name, branch)."""
    out = {}
    wb = load_workbook(PORTFOLIO_FILE, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r[1]:
            continue
        name = str(r[1]).strip()
        branch = str(r[5] or "").replace("BRANCH", "").strip().title() or None
        out[digits(r[3])] = (name, branch)
    wb.close()
    return out


def wipe():
    with engine.begin() as conn:
        for table in reversed(Base.metadata.sorted_tables):
            conn.execute(table.delete())


def main():
    Base.metadata.create_all(bind=engine)
    portfolio = load_portfolio()
    db = SessionLocal()
    used_emails = set()
    counters = {"TC": 0, "FO": 0, "MG": 0}

    def email(name):
        base = slugify(name)
        e = f"{base}@ssd.local"
        i = 2
        while e in used_emails:
            e = f"{base}{i}@ssd.local"
            i += 1
        used_emails.add(e)
        return e

    def code(prefix):
        counters[prefix] += 1
        return f"{prefix}{counters[prefix]:03d}"

    def add(name, role, branch, **kw):
        u = models.User(name=name, email=email(name), role=role, branch=branch,
                        banks=[BANK], hashed_password=PW, is_active=True, **kw)
        db.add(u)
        db.flush()
        return u

    try:
        wipe()

        admin = models.User(name="Admin", email="admin@ssd.local", role="admin",
                            emp_code="AD001", banks=[BANK], hashed_password=PW, is_active=True)
        db.add(admin)
        db.flush()

        # 5 branches, each with a manager; Visakhapatnam is HQ.
        for b in BRANCHES:
            add(f"{b} Manager", "manager", b, emp_code=code("MG"))
        # head-office (cross-branch) sits at the HQ branch
        db.add(models.User(name="Head Office", email="ho@ssd.local", role="headoffice",
                           branch=HEAD_BRANCH, emp_code="HO001", banks=[BANK],
                           hashed_password=PW, is_active=True))
        db.flush()

        # read the FR sheet
        with open(FR_FILE, "rb") as fh:
            records, _ = import_workbook(fh.read(), default_bank=BANK, sheet_name="MAIN")
        kw_list = [record_to_case_kwargs(r) for r in records]

        def clean(v):
            return (v or "").strip()

        def is_real_fos(t):
            return bool(t) and any(ch.isdigit() for ch in t)

        # --- FOS accounts, mapped to their branch from the portfolio ---
        fos_texts = sorted({clean(k.get("fos_name")) for k in kw_list if is_real_fos(clean(k.get("fos_name")))})
        fos_by_text = {}
        for txt in fos_texts:
            ph = digits(txt)
            if ph in portfolio:
                name, branch = portfolio[ph]
            else:
                name = txt.split(",")[0].split("/")[0].strip().title()
                area = txt.split("/")[1].split(",")[0].strip().upper() if "/" in txt else ""
                branch = AREA_BRANCH.get(area)
            u = add(name, "fos", branch, emp_code=code("FO"), assigned_products=[PRODUCT])
            fos_by_text[txt] = u

        # --- Telecaller accounts, all at Visakhapatnam ---
        caller_names = sorted({clean(k.get("caller_name")) for k in kw_list
                               if clean(k.get("caller_name")) and clean(k.get("caller_name")).upper() != "REMOVE"})
        caller_by_name = {}
        for nm in caller_names:
            u = add(nm.title(), "telecaller", HEAD_BRANCH, emp_code=code("TC"))
            caller_by_name[nm.upper()] = u

        # --- import cases ---
        batch = models.ImportBatch(filename="fr_xbkt_july26.xlsx", bank=BANK, sheet="MAIN",
                                   rows_total=len(kw_list), uploaded_by=admin.id)
        db.add(batch)
        db.flush()

        imported = 0
        for kw in kw_list:
            if not kw.get("customer_name") and not kw.get("account_no"):
                continue
            kw = dict(kw)
            kw.update(bank=BANK, product=PRODUCT, segment=SEGMENT, import_batch_id=batch.id)
            cu = caller_by_name.get(clean(kw.get("caller_name")).upper())
            if cu:
                kw["assigned_caller_id"] = cu.id
                kw["caller_name"] = cu.name
            fu = fos_by_text.get(clean(kw.get("fos_name")))
            if fu:
                kw["assigned_fos_id"] = fu.id
                kw["branch"] = fu.branch          # case's branch = its FOS's branch
            db.add(models.Case(**kw))
            imported += 1
        batch.rows_imported = imported
        db.commit()

        # summary
        from collections import Counter
        by_branch = Counter(u.branch for u in db.query(models.User).filter(models.User.role == "fos").all())
        print("=== Seed complete ===")
        print(f"  branches:    {', '.join(BRANCHES)}   (HQ = {HEAD_BRANCH})")
        print(f"  admin: admin@ssd.local | head office: ho@ssd.local | all pw: Test@1234")
        print(f"  managers:    {counters['MG']}  (e.g. hyderabad.manager@ssd.local)")
        print(f"  FOS:         {counters['FO']}  mapped by branch -> {dict(by_branch)}")
        print(f"  telecallers: {counters['TC']}  (all at {HEAD_BRANCH})")
        print(f"  cases:       {imported}  ({BANK} · {PRODUCT} · {SEGMENT})")
    finally:
        db.close()


if __name__ == "__main__":
    main()
